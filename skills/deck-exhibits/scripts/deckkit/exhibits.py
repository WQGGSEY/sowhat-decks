"""Exhibits as native PowerPoint objects: charts, tables and shapes.

Style rules (deck-exhibits/references/chart-style.md): direct labels instead of
legends, one accent for the point of the slide and greys for the rest, units in
the exhibit title, no gridline clutter, labels at >= 12 pt.
"""
from __future__ import annotations

import itertools
import math

from lxml import etree
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Pt

from . import spec as deckspec, text_fit, theme
from .draw import Ctx, common_size, connector, rect, text
from .grid import Box, pt, to_pt
from .numbers import auto_format, fmt_number, format_cell
from .text import P, fit_box, write

LABEL_PT = 12
END_LABEL_ONE_LINE = 0.18  # widest one-line line-chart end label, as a share of the chart width
NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # "No Style, No Grid"


# ------------------------------------------------------------------ helpers

def exhibit_title(ex: dict) -> str:
    title, unit = ex["title"], ex.get("unit")
    if unit and unit.lower() not in title.lower():
        title = f"{title}, {unit}"
    return title


def nice_scale(lo: float, hi: float, ticks: int = 5) -> tuple[float, float, float]:
    if hi == lo:
        hi = lo + 1
    raw = (hi - lo) / ticks
    mag = 10 ** math.floor(math.log10(raw))
    step = next(m * mag for m in (1, 2, 2.5, 5, 10) if m * mag >= raw)
    return math.floor(lo / step) * step, math.ceil(hi / step) * step, step


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def header(ctx: Ctx, slide, title: str, box: Box) -> Box:
    """Exhibit title line (what + unit). Returns the area left for the exhibit."""
    paras = [P(title, bold=True)]
    fit = fit_box(paras, box, ctx.frame, max_size=14, min_size=12, max_lines=2)
    h = pt(fit.height_pt) + ctx.s(0.04)
    text(ctx, slide, Box(box.x, box.y, box.w, h), paras, role="exhibit-title", size=fit.size, max_lines=2)
    return Box(box.x, box.y + h + ctx.s(0.12), box.w, max(0, box.h - h - ctx.s(0.12)))


def _plain_chart(chart) -> None:
    """No legend, no title, no border, 12 pt ink text."""
    chart.has_legend = False
    chart.has_title = False
    chart.font.size = Pt(LABEL_PT)
    chart.font.color.rgb = _rgb(theme.INK)
    space = chart._chartSpace
    old = space.find(qn("c:spPr"))
    if old is not None:
        space.remove(old)
    sppr = parse_xml(f'<c:spPr {nsdecls("c", "a")}><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>')
    space.find(qn("c:chart")).addnext(sppr)


def _plot_layout(chart, x: float, y: float, w: float, h: float) -> None:
    """Fix the inner plot area (fractions of the chart frame) so label positions are known."""
    plot_area = chart._chartSpace.find(qn("c:chart")).find(qn("c:plotArea"))
    old = plot_area.find(qn("c:layout"))
    if old is not None:
        plot_area.remove(old)
    layout = parse_xml(
        f'<c:layout {nsdecls("c")}><c:manualLayout>'
        '<c:layoutTarget val="inner"/><c:xMode val="edge"/><c:yMode val="edge"/>'
        f'<c:x val="{x:.4f}"/><c:y val="{y:.4f}"/><c:w val="{w:.4f}"/><c:h val="{h:.4f}"/>'
        '</c:manualLayout></c:layout>')
    plot_area.insert(0, layout)


def _axis_text(axis, color: str = theme.INK) -> None:
    axis.tick_labels.font.size = Pt(LABEL_PT)
    axis.tick_labels.font.color.rgb = _rgb(color)
    axis.major_tick_mark = XL_TICK_MARK.NONE
    axis.minor_tick_mark = XL_TICK_MARK.NONE


def _category_axis(chart) -> None:
    ca = chart.category_axis
    _axis_text(ca)
    ca.has_major_gridlines = False
    ca.format.line.color.rgb = _rgb(theme.GREY_2)
    ca.format.line.width = Pt(0.75)
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW


def _hide_value_axis(chart, lo: float | None = None, hi: float | None = None) -> None:
    va = chart.value_axis
    va.has_major_gridlines = False
    va.visible = False
    if lo is not None:
        va.minimum_scale = lo
    if hi is not None:
        va.maximum_scale = hi


def _fill(fmt, color: str) -> None:
    fmt.fill.solid()
    fmt.fill.fore_color.rgb = _rgb(color)


def _label_offset(dlbl, dx: float, dy: float) -> None:
    """Move one data label by a fraction of the chart size."""
    idx = dlbl.find(qn("c:idx"))
    old = dlbl.find(qn("c:layout"))
    if old is not None:
        dlbl.remove(old)
    layout = parse_xml(
        f'<c:layout {nsdecls("c")}><c:manualLayout>'
        f'<c:x val="{dx:.4f}"/><c:y val="{dy:.4f}"/></c:manualLayout></c:layout>')
    idx.addnext(layout)


def _hide_label(series, idx: int) -> None:
    dlbl = series._element.get_or_add_dLbls().get_or_add_dLbl_for_point(idx)
    for child in list(dlbl):
        if child.tag != qn("c:idx"):
            dlbl.remove(child)
    etree.SubElement(dlbl, qn("c:delete"), val="1")


def _spread(ys: list[float], gap: float, lo: float, hi: float) -> list[float]:
    """Push label centres apart so they are at least gap apart, staying within [lo, hi]."""
    order = sorted(range(len(ys)), key=lambda i: ys[i])
    out = list(ys)
    for a, b in zip(order, order[1:]):
        out[b] = max(out[b], out[a] + gap)
    overflow = out[order[-1]] - hi if order else 0
    if overflow > 0:
        for i in order:
            out[i] -= overflow
        for a, b in zip(reversed(order[:-1]), reversed(order[1:])):
            out[a] = min(out[a], out[b] - gap)
    return [max(lo, y) for y in out]


def _label_room(values: list, k: float) -> tuple[float, float]:
    """Axis (min, max) that keeps a share k of the plot free past each bar end, on the side it points."""
    top, bottom = max(0, max(values, default=0)), min(0, min(values, default=0))
    if top == bottom:
        return 0, 1
    span = (top - bottom) / (1 - k * (top > 0) - k * (bottom < 0))
    return bottom - (k * span if bottom < 0 else 0), top + (k * span if top > 0 else 0)


def _chart_frame(slide, kind, data, box: Box):
    gf = slide.shapes.add_chart(kind, box.x, box.y, box.w, box.h, data)
    gf.name = "sw:chart"
    return gf.chart


# ------------------------------------------------------------------ bar

def bar(ctx: Ctx, slide, ex: dict, box: Box) -> None:
    cats, vals = list(ex["categories"]), list(ex["values"])
    hi_set = {cats[h] if isinstance(h, int) else h for h in ex.get("highlight", [])}
    if ex.get("sort", "none") != "none":
        sign = -1 if ex["sort"] == "desc" else 1  # missing values always go last
        pairs = sorted(zip(cats, vals), key=lambda p: (p[1] is None, sign * (p[1] or 0)))
        cats, vals = [p[0] for p in pairs], [p[1] for p in pairs]
    fmt = ex.get("number_format") or auto_format(vals)
    data = CategoryChartData(number_format=fmt)
    data.categories = cats
    data.add_series(ex["title"][:30], vals)
    horizontal = ex.get("orientation", "horizontal") == "horizontal"
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    chart = _chart_frame(slide, kind, data, box)
    _plain_chart(chart)
    plot = chart.plots[0]
    plot.gap_width = 110 if len(cats) <= 5 else 80 if len(cats) <= 8 else 60
    plot.vary_by_categories = False
    series = plot.series[0]
    _fill(series.format, theme.GREY_2)
    series.format.line.fill.background()
    labels = series.data_labels  # series level, so per-point tweaks below keep the others
    labels.number_format = fmt
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.size = Pt(LABEL_PT)
    labels.font.color.rgb = _rgb(theme.INK)
    labels.show_value = True
    for i, c in enumerate(cats):
        if c in hi_set:
            _fill(series.points[i].format, ctx.accent)
            dl = series.points[i].data_label
            dl.font.bold = True
            dl.font.size = Pt(LABEL_PT)
            dl.font.color.rgb = _rgb(ctx.accent)
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
    # A fixed plot area and an axis range with room for the value labels: every label, past the end of
    # its bar (below it when the value is negative), stays inside the plot, so it can never run into the
    # category labels drawn outside the plot.
    nums = [v for v in vals if v is not None]
    font = ctx.frame.body_font
    w_pt, h_pt = to_pt(box.w), to_pt(box.h)
    if horizontal:
        x = min(0.45, (max(text_fit.text_width(c, font, LABEL_PT) for c in cats) + 10) / w_pt)
        y, pw, ph = 0.02, 1 - x - 0.01, 0.96
        label = max((text_fit.text_width(fmt_number(v, fmt), font, LABEL_PT, bold=True) for v in nums), default=0)
        k = (label * text_fit.SAFETY + 6) / (pw * w_pt)
    else:
        slot = 0.96 * w_pt / len(cats)
        lines = 2 if any(text_fit.text_width(c, font, LABEL_PT) > slot - 4 for c in cats) else 1
        x, y, pw = 0.02, 0.04, 0.96
        ph = 1 - y - (lines * LABEL_PT * 1.2 + 8) / h_pt
        k = (LABEL_PT * 1.2 + 4) / (ph * h_pt)
    _plot_layout(chart, x, y, pw, ph)
    _hide_value_axis(chart, *_label_room(nums, min(k, 0.3)))
    _category_axis(chart)
    if horizontal:
        chart.category_axis.reverse_order = True


# ------------------------------------------------------------------ line

def line(ctx: Ctx, slide, ex: dict, box: Box) -> None:
    cats, series_spec = ex["categories"], ex["series"]
    all_vals = [v for s in series_spec for v in s["values"] if v is not None]
    fmt = ex.get("number_format") or auto_format(all_vals)
    data = CategoryChartData(number_format=fmt)
    data.categories = cats
    for s in series_spec:
        data.add_series(s["name"], s["values"])
    chart = _chart_frame(slide, XL_CHART_TYPE.LINE, data, box)
    _plain_chart(chart)

    lo_v, hi_v = min(all_vals, default=0), max(all_vals, default=1)
    lo, hi, step = nice_scale(0 if 0 <= lo_v <= 0.5 * hi_v else lo_v, hi_v)
    va = chart.value_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = lo, hi, step
    _axis_text(va, theme.INK_2)
    va.tick_labels.number_format = fmt
    va.tick_labels.number_format_is_linked = False
    va.format.line.fill.background()
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = _rgb(theme.GREY_3)
    va.major_gridlines.format.line.width = Pt(0.5)
    _category_axis(chart)

    hl = ex.get("highlight")
    others = itertools.cycle(theme.GREY_RAMP[:4])
    colors = [ctx.accent if s["name"] == hl or (hl is None and len(series_spec) == 1) else next(others)
              for s in series_spec]
    series = list(chart.plots[0].series)  # same order as the spec
    for ser, color in zip(series, colors):
        ser.smooth = False
        ser.marker.style = XL_MARKER_STYLE.NONE
        ser.format.line.color.rgb = _rgb(color)
        ser.format.line.width = Pt(2.5 if color == ctx.accent else 1.75)

    # End labels "Series value" at each series' last point; the plot leaves room for them on the right.
    # A label wider than END_LABEL_ONE_LINE of the chart is written as two lines, name then value, so a renderer
    # can only break between them (LibreOffice wraps data labels at about a fifth of the chart width
    # and would otherwise split the number). wrap="none" asks renderers not to re-wrap at all.
    font = ctx.frame.body_font
    w_pt, h_pt = to_pt(box.w), to_pt(box.h)
    ends = []
    for i, s in enumerate(series_spec):
        last = max((j for j, v in enumerate(s["values"]) if v is not None), default=None)
        if last is not None:
            name, value = s["name"], fmt_number(s["values"][last], fmt)
            ends.append((i, last, [name, value]))
    width = lambda t, i: text_fit.text_width(t, font, LABEL_PT, bold=colors[i] == ctx.accent)  # noqa: E731
    two_lines = any(width(" ".join(lines), i) > END_LABEL_ONE_LINE * w_pt for i, _, lines in ends)
    ends = [(i, last, lines if two_lines else [" ".join(lines)]) for i, last, lines in ends]
    label_w = max((width(t, i) for i, _, lines in ends for t in lines), default=0)
    axis_w = max(text_fit.text_width(fmt_number(v, fmt), font, LABEL_PT) for v in (lo, hi)) + 8
    x, y = min(0.25, axis_w / w_pt), 0.04
    pw = 1 - x - min(0.45, (label_w + 14) / w_pt)
    ph = 1 - y - (LABEL_PT * 1.6 + 6) / h_pt
    _plot_layout(chart, x, y, pw, ph)
    gap = (2 if two_lines else 1) * LABEL_PT * 1.25 / h_pt
    want = [y + ph * (hi - series_spec[i]["values"][last]) / (hi - lo) for i, last, _ in ends]
    for (i, last, lines), y0, y1 in zip(ends, want, _spread(want, gap, 0.0, 1 - gap)):
        dl = series[i].points[last].data_label
        dl.has_text_frame = True
        dl.text_frame.text = "\n".join(lines)  # one paragraph per line
        dl.text_frame._txBody.find(qn("a:bodyPr")).set("wrap", "none")
        for paragraph in dl.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.LEFT  # name and value start at the same x, next to the line end
            for run in paragraph.runs:
                run.font.size = Pt(LABEL_PT)
                run.font.bold = colors[i] == ctx.accent
                run.font.color.rgb = _rgb(colors[i])
        dl.position = XL_LABEL_POSITION.RIGHT
        if abs(y1 - y0) > 1e-4:  # nudge labels that would collide
            _label_offset(series[i]._element.get_or_add_dLbls().get_or_add_dLbl_for_point(last), 0, y1 - y0)


# ------------------------------------------------------------------ stacked bars

def stacked(ctx: Ctx, slide, ex: dict, box: Box) -> None:
    share = ex["type"] == "stacked_bar_100"
    cats, series_spec = ex["categories"], ex["series"]
    n = len(cats)
    totals = [sum((s["values"][i] or 0) for s in series_spec) for i in range(n)]
    if share:
        rows = [[(s["values"][i] or 0) / totals[i] if totals[i] else 0 for i in range(n)] for s in series_spec]
        fmt = ex.get("number_format") or "0%"
    else:
        rows = [list(s["values"]) for s in series_spec]
        fmt = ex.get("number_format") or auto_format([v for r in rows for v in r if v is not None])
    data = CategoryChartData(number_format=fmt)
    data.categories = cats
    for s, r in zip(series_spec, rows):
        data.add_series(s["name"], r)
    kind = XL_CHART_TYPE.COLUMN_STACKED_100 if share else XL_CHART_TYPE.COLUMN_STACKED
    chart = _chart_frame(slide, kind, data, box)
    _plain_chart(chart)
    plot = chart.plots[0]
    plot.gap_width = 55
    plot.overlap = 100

    hi = 1.0 if share else nice_scale(0, max(totals) if totals else 1)[1]
    _hide_value_axis(chart, lo=0, hi=hi)
    _category_axis(chart)

    font = ctx.frame.body_font
    w_pt, h_pt = to_pt(box.w), to_pt(box.h)
    hl = ex.get("highlight")
    label_w = max(text_fit.text_width(s["name"], font, LABEL_PT, bold=s["name"] == hl) for s in series_spec)
    show_totals = not share and ex.get("show_totals", True)
    x, right = 0.01, min(0.4, (label_w + 16) / w_pt)
    y = (LABEL_PT * 1.6) / h_pt if show_totals else 0.02
    bottom = (LABEL_PT * 1.6 + 6) / h_pt
    pw, ph = 1 - x - right, 1 - y - bottom
    _plot_layout(chart, x, y, pw, ph)

    greys = iter(theme.GREY_RAMP[1:] if hl else theme.GREY_RAMP)  # with a highlight, start lighter
    colors = {s["name"]: (ctx.accent if s["name"] == hl else next(greys)) for s in series_spec}
    light = {theme.GREY_1, theme.GREY_2, theme.GREY_3}
    slot = pw * w_pt / n
    bar_w = slot / 1.55
    for s_i, (s, ser) in enumerate(zip(series_spec, plot.series)):
        color = colors[s["name"]]
        _fill(ser.format, color)
        ser.format.line.color.rgb = _rgb(theme.WHITE)
        ser.format.line.width = Pt(1)
        ser.data_labels.show_value = True
        ser.data_labels.number_format = fmt
        ser.data_labels.number_format_is_linked = False
        ser.data_labels.position = XL_LABEL_POSITION.CENTER
        ser.data_labels.font.size = Pt(LABEL_PT)
        ser.data_labels.font.bold = color == ctx.accent
        ser.data_labels.font.color.rgb = _rgb(theme.INK if color in light else theme.WHITE)
        for i in range(n):  # hide labels that do not fit inside their segment
            v = rows[s_i][i] or 0
            seg_h = ph * h_pt * v / hi
            label = fmt_number(v, fmt)
            if seg_h < LABEL_PT * 1.25 or text_fit.text_width(label, font, LABEL_PT) > bar_w - 4:
                _hide_label(ser, i)

    # direct series labels to the right of the last column, totals above columns
    plot_x, plot_y = box.x + pt(x * w_pt), box.y + pt(y * h_pt)
    last_right = plot_x + pt(slot * (n - 1) + (slot + bar_w) / 2)
    mids, cum = [], 0.0
    for s_i in range(len(series_spec)):
        v = rows[s_i][-1] or 0
        mids.append(1 - (cum + v / 2) / hi)
        cum += v
    gap = LABEL_PT * 1.3 / (ph * h_pt)
    placed = _spread(mids, gap, 0, 1 - gap / 2)
    lab_h = pt(LABEL_PT * 1.3)
    for s, m in zip(series_spec, placed):
        cy = plot_y + pt(m * ph * h_pt)
        lbox = Box(last_right + pt(6), cy - lab_h // 2, box.right - last_right - pt(6), lab_h)
        color = colors[s["name"]]
        text(ctx, slide, lbox, [P(s["name"], bold=color == ctx.accent)], role="label", size=LABEL_PT,
             color=ctx.accent if color == ctx.accent else theme.INK_2, anchor="m", max_lines=1)
    if show_totals:
        for i, t in enumerate(totals):
            cx = plot_x + pt(slot * i + slot / 2)
            top_y = plot_y + pt(ph * h_pt * (1 - t / hi))
            tw = pt(slot)
            text(ctx, slide, Box(cx - tw // 2, top_y - lab_h - pt(2), tw, lab_h), [P(fmt_number(t, fmt), bold=True)],
                 role="label", size=LABEL_PT, align="c", anchor="b", max_lines=1)


# ------------------------------------------------------------------ tables

def _cell_borders(cell, bottom: tuple[str, float] | None, fill: str | None) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    del tcpr[:]
    for side in ("lnL", "lnR", "lnT", "lnB"):
        if side == "lnB" and bottom:
            ln = etree.SubElement(tcpr, qn(f"a:{side}"), w=str(int(Pt(bottom[1]))), cap="flat", cmpd="sng")
            sf = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr"), val=bottom[0])
        else:
            ln = etree.SubElement(tcpr, qn(f"a:{side}"), w="0")
            etree.SubElement(ln, qn("a:noFill"))
    if fill:
        sf = etree.SubElement(tcpr, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val=fill)
    else:
        etree.SubElement(tcpr, qn("a:noFill"))


def table(ctx: Ctx, slide, t: dict, box: Box) -> None:
    """Native table: bold header on a rule, hairlines between rows, numbers right-aligned."""
    cols, rows = t["columns"], t["rows"]
    hl = t.get("highlight") or {}
    lit = {(r, c) for r in range(len(rows)) for c in range(len(cols))
           if r in hl.get("rows", []) or c in hl.get("cols", []) or [r, c] in hl.get("cells", [])}
    numeric = [any(isinstance(r[c], (int, float)) for r in rows)
               and all(isinstance(r[c], (int, float)) or r[c] is None for r in rows) for c in range(len(cols))]
    align = t.get("align") or ["right" if numeric[c] and c else "left" for c in range(len(cols))]
    font = ctx.frame.body_font
    pad_x, pad_y = ctx.s(0.1), ctx.s(0.05)
    # one paragraph per cell, bold exactly as it will be written (header and highlighted cells)
    cells = [[P(c, bold=True, font=font) for c in cols]] + [
        [P(format_cell(v, t.get("number_format")), bold=(r, c) in lit, font=font) for c, v in enumerate(row)]
        for r, row in enumerate(rows)]

    need = [min(max(max(text_fit.text_width(p.text, font, 14, p.bold) for p in col) + 2 * to_pt(pad_x), 60), 320)
            for col in zip(*cells)]
    widths = [int(box.w * w / sum(need)) for w in need]
    widths[-1] = box.w - sum(widths[:-1])

    def heights(size: float) -> list[int]:
        return [pt(max(text_fit.measure([p], size, to_pt(w - 2 * pad_x))[2] for p, w in zip(row, widths)))
                + 2 * pad_y + pt(2) for row in cells]

    size, hs = 14, heights(14)
    while size > 12 and sum(hs) > box.h:
        size -= 1
        hs = heights(size)
    gf = slide.shapes.add_table(len(cells), len(cols), box.x, box.y, box.w, min(sum(hs), box.h))
    gf.name = "sw:table"
    tbl = gf.table
    tbl.first_row, tbl.horz_banding = True, False
    gf._element.graphic.graphicData.tbl.tblPr.find(qn("a:tableStyleId")).text = NO_STYLE_TABLE
    for c, w in enumerate(widths):
        tbl.columns[c].width = w
    for r, h in enumerate(hs):
        tbl.rows[r].height = h
    for r, row in enumerate(cells):
        for c, p in enumerate(row):
            cell = tbl.cell(r, c)
            rule = (theme.GREY_2, 0.75) if r == len(rows) else (theme.INK, 1.0) if r == 0 else (theme.GREY_3, 0.75)
            _cell_borders(cell, rule, ctx.accent_tint if (r - 1, c) in lit else None)
            cell.margin_left = cell.margin_right = pad_x
            cell.margin_top = cell.margin_bottom = pad_y
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE if r else MSO_ANCHOR.BOTTOM
            write(cell.text_frame, [p], [size], align=align[c][0] if c < len(align) else "l", anchor=None,
                  insets=None, lang=ctx.lang, default_color=theme.INK_2 if r == 0 else theme.INK)
    fit = text_fit.Fit(size, [size], [len(hs)], to_pt(sum(hs)), sum(hs) > box.h)
    ctx.report.add(ctx.slide_no, "table", fit, " | ".join(cols))


# ------------------------------------------------------------------ 2x2 matrix

def matrix_2x2(ctx: Ctx, slide, ex: dict, box: Box) -> None:
    gutter = ctx.s(1.25)
    foot = ctx.s(0.75)
    area = Box(box.x + gutter, box.y, box.w - gutter, box.h - foot)
    gap = ctx.s(0.05)
    quads = [q for row in area.rows(2, gap) for q in row.cols(2, gap)]  # TL, TR, BL, BR
    hq = ex.get("highlight_quadrant")
    # Quadrant names sit on the outer edges (top row at the top, bottom row at the bottom), so the
    # middle of the matrix stays free for items. Each name box spans its quadrant edge to edge and
    # keeps the padding as text insets, so its edges line up with the quadrant's. name_h is the
    # strip a name takes, padding included.
    pad, side = ctx.s(0.1), ctx.s(0.12)
    names = [[P(label, bold=True)] for label in ex["quadrants"]]
    probe = [(q.inset(side, 0, side, 0).moved(h=ctx.s(0.5)), p) for q, p in zip(quads, names)]
    name_size = common_size(ctx, probe, max_size=14, min_size=12, max_lines=2)
    name_h = pad + ctx.s(0.04) + max(pt(fit_box(p, b, ctx.frame, max_size=name_size, min_size=name_size).height_pt)
                                     for b, p in probe)
    for i, (q, paras) in enumerate(zip(quads, names)):
        lit = i == hq
        rect(slide, q, fill=ctx.accent_tint if lit else theme.GREY_4, name="sw:quadrant")
        if paras[0].text:
            top_row = i < 2
            text(ctx, slide, q.moved(y=q.y if top_row else q.bottom - name_h, h=name_h), paras, role="label",
                 size=name_size, max_lines=2, anchor="t" if top_row else "b",
                 insets=(side, pad, side, 0) if top_row else (side, 0, side, pad),
                 color=ctx.accent if lit else theme.INK_2)
    for x2, y2 in ((area.x, area.y), (area.right, area.bottom)):
        connector(slide, area.x, area.bottom, x2, y2, color=theme.INK_2, width=1.25, name="sw:axis", arrow=True)
    lab_h = ctx.s(0.3)
    text(ctx, slide, Box(box.x, area.y + area.h // 2 - ctx.s(0.5), gutter - ctx.s(0.12), ctx.s(1.0)),
         [P(ex["y_label"], bold=True)], role="label", size=12, align="r", anchor="m", max_lines=3)
    if ex.get("y_high"):
        text(ctx, slide, Box(box.x, area.y, gutter - ctx.s(0.12), lab_h), [P(ex["y_high"])], role="label",
             size=12, align="r", color=theme.INK_2, max_lines=1)
    if ex.get("y_low"):
        text(ctx, slide, Box(box.x, area.bottom - lab_h, gutter - ctx.s(0.12), lab_h), [P(ex["y_low"])],
             role="label", size=12, align="r", anchor="b", color=theme.INK_2, max_lines=1)
    row = Box(area.x, area.bottom + ctx.s(0.08), area.w, lab_h)
    third = row.w // 3
    if ex.get("x_low"):
        text(ctx, slide, row.moved(w=third), [P(ex["x_low"])], role="label", size=12, color=theme.INK_2, max_lines=1)
    if ex.get("x_high"):
        text(ctx, slide, row.moved(x=row.right - third, w=third), [P(ex["x_high"])], role="label", size=12,
             align="r", color=theme.INK_2, max_lines=1)
    text(ctx, slide, Box(area.x + third, row.y, third, lab_h * 2), [P(ex["x_label"], bold=True)], role="label",
         size=12, align="c", max_lines=2)

    dot = ctx.s(0.18)
    lh = ctx.s(0.28)
    font = ctx.frame.body_font
    keep_out = name_h + ctx.s(0.04) + lh // 2  # item labels stay clear of the quadrant names
    for item in ex.get("items", []):
        cx = area.x + int(item["x"] * area.w)
        cy = area.bottom - int(item["y"] * area.h)
        cx = min(max(cx, area.x + dot), area.right - dot)
        cy = min(max(cy, area.y + keep_out), area.bottom - keep_out)
        color = ctx.accent if item.get("highlight") else theme.INK_2
        rect(slide, Box(cx - dot // 2, cy - dot // 2, dot, dot), fill=color, shape=MSO_SHAPE.OVAL, name="sw:dot")
        lw = pt(text_fit.text_width(item["label"], font, 12, bold=item.get("highlight", False)) * 1.08) + ctx.s(0.05)
        left = cx + dot if cx + dot + lw <= area.right else cx - dot - lw
        top = min(max(cy - lh // 2, area.y), area.bottom - lh)
        text(ctx, slide, Box(max(left, area.x), top, lw, lh), [P(item["label"], bold=item.get("highlight", False))],
             role="label", size=12, color=color, anchor="m", max_lines=1)


# ------------------------------------------------------------------ process chevrons

def process(ctx: Ctx, slide, ex: dict, box: Box) -> None:
    """Chevrons with the label in a text box over each chevron body, grouped so they move together."""
    steps = ex["steps"]
    n = len(steps)
    gap = ctx.s(0.06)
    chev_h = min(ctx.s(0.95), box.h // 3)
    w = (box.w - gap * (n - 1)) // n
    notch = int(chev_h * 0.3)
    pad = ctx.s(0.06)
    hl = ex.get("highlight")

    def label_box(i: int) -> Box:
        x = box.x + i * (w + gap)
        left = x + (pad if i == 0 else notch + pad // 2)
        return Box(left, box.y + pad // 2, x + w - notch - pad // 2 - left, chev_h - pad)

    labels = [(label_box(i), [P(s["label"], bold=True)]) for i, s in enumerate(steps)]
    label_size = common_size(ctx, labels, max_size=16, min_size=12, max_lines=3)
    top = box.y + chev_h + ctx.s(0.25)
    details = [(Box(box.x + i * (w + gap) + (0 if i == 0 else notch // 2), top, w - notch, box.bottom - top),
                [P(s.get("detail", ""))]) for i, s in enumerate(steps)]
    detail_size = common_size(ctx, details, max_size=16, min_size=12)
    for i, s in enumerate(steps):
        lit = i == hl
        group = slide.shapes.add_group_shape()
        group.name = "sw:step"
        sp = rect(group, Box(box.x + i * (w + gap), box.y, w, chev_h), fill=ctx.accent if lit else theme.GREY_4,
                  shape=MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON, name="sw:step-shape")
        sp.adjustments[0] = notch / min(w, chev_h)
        lbox, paras = labels[i]
        text(ctx, group, lbox, paras, size=label_size, align="c", anchor="m", max_lines=3,
             color=theme.WHITE if lit else theme.INK)
        if s.get("detail"):
            dbox, dparas = details[i]
            text(ctx, slide, dbox, dparas, size=detail_size, color=theme.INK_2)


# ------------------------------------------------------------------ dispatch

RENDERERS = {"bar": bar, "line": line, "stacked_bar": stacked, "stacked_bar_100": stacked,
             "highlight_table": table, "matrix_2x2": matrix_2x2, "process": process}


def register(schema: dict, render, check=None) -> None:
    """Add an exhibit type from an add-on skill; the validator and every exhibit layout accept it.

    schema: the type's JSON Schema branch, with properties.type.const = the type name. It may $ref
            deck-spec.schema.json's $defs, e.g. {"$ref": "#/$defs/ex_common/title"}.
    render: render(ctx, slide, ex, box) draws the exhibit in box, under the shared exhibit title.
            Its return value is ignored (it may return None): layouts use the area that
            exhibits.render() returns.
    check:  optional check(ex, path, errors) for rules a schema cannot express. Like the built-in
            checks it runs only on a schema-valid spec and appends readable errors.
    Registering a name again replaces it. Built-in types cannot be replaced (ValueError).
    """
    RENDERERS[deckspec.add_exhibit(schema, check)] = render


def render(ctx: Ctx, slide, ex: dict, box: Box) -> Box:
    """Exhibit title, then the exhibit. Returns the box the exhibit itself fills (under its title)."""
    area = header(ctx, slide, exhibit_title(ex), box)
    RENDERERS[ex["type"]](ctx, slide, ex, area)
    return area
