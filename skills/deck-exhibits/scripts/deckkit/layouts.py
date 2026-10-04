"""The 12 layouts (plus ghost). Each draws one slide's body on the 12-column grid."""
from __future__ import annotations

import re

from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER

from . import exhibits, theme
from .draw import Ctx, common_size, fill_placeholder, hline, keep_placeholders, rect, set_title, text
from .grid import Box, pt
from .template_map import TITLE_TYPES, box_of, placeholders_by_type
from .text import P, fit_box


# "[DATA NEEDED]", "[DATA NEEDED: what]" (deck-storyline's form), any case and spacing
DATA_NEEDED = re.compile(r"\[\s*DATA\s+NEEDED\s*(?:[:：][^\]]*)?\]", re.IGNORECASE)


def _head_detail(head: str, detail: str | None, scale: float) -> list[P]:
    """Bold lead line, optional grey detail under it at `scale` of the lead's size."""
    paras = [P(head, bold=True)]
    if detail:
        paras.append(P(detail, scale=scale, min_size=12, color=theme.INK_2, space_before=0.3))
    return paras


# ------------------------------------------------------------------ cover / section

def cover(ctx: Ctx, slide, s: dict) -> None:
    f, meta = ctx.frame, ctx.meta
    subtitle = s.get("subtitle", meta.get("subtitle", ""))
    by = " · ".join(x for x in (s.get("author") or meta.get("author") or meta.get("organization"),
                                 s.get("date") or meta.get("date")) if x)
    lines = ([P(subtitle)] if subtitle else []) + ([P(by, scale=0.7, min_size=12, space_before=0.6)] if by else [])
    phs = placeholders_by_type(slide)
    sub_ph = phs.get(PP_PLACEHOLDER.SUBTITLE) or phs.get(PP_PLACEHOLDER.BODY)
    if not lines:
        sub_ph = None
    keep_placeholders(slide, TITLE_TYPES + ((sub_ph.placeholder_format.type,) if sub_ph is not None else ()))
    title_ph, _ = set_title(ctx, slide, s.get("title") or meta["title"], role="cover-title", max_size=40,
                            min_size=28, max_lines=3)
    if sub_ph is not None:
        fill_placeholder(ctx, sub_ph, lines, role="body", max_size=20, min_size=14)
    elif lines:
        text(ctx, slide, Box(f.title.x, f.footer_bottom - ctx.s(1.4), f.title.w, ctx.s(1.0)), lines,
             max_size=20, min_size=14, color=theme.INK_2)
    if f.default:  # short accent rule above the title, part of the built-in look only
        rect(slide, Box(f.title.x, int(title_ph.top) - ctx.s(0.25), ctx.s(0.9), ctx.s(0.06)),
             fill=ctx.accent, name="sw:accent-rule")


def section(ctx: Ctx, slide, s: dict) -> None:
    f = ctx.frame
    subtitle = s.get("subtitle", "")
    body_ph = placeholders_by_type(slide).get(PP_PLACEHOLDER.BODY) if subtitle else None
    keep_placeholders(slide, TITLE_TYPES + ((PP_PLACEHOLDER.BODY,) if body_ph is not None else ()))
    number = s.get("number")
    title = s["title"] if number is None or f.default else f"{number}  {s['title']}"
    title_ph, _ = set_title(ctx, slide, title, role="section-title", max_size=36, min_size=28, max_lines=2)
    t = box_of(title_ph)
    if body_ph is not None:
        fill_placeholder(ctx, body_ph, [P(subtitle)], role="body", max_size=18, min_size=14)
    elif subtitle:
        text(ctx, slide, Box(t.x, t.bottom + ctx.s(0.2), t.w, ctx.s(1.0)), [P(subtitle)], max_size=18, min_size=14,
             color=theme.INK_2)
    if number is not None and f.default:  # big accent number left of the title, built-in look only
        grid = ctx.grid(f.body)
        text(ctx, slide, Box(grid.x(1), t.y, grid.width(2), t.h), [P(str(number), bold=True)], role="number",
             size=54, color=ctx.accent, anchor="b", max_lines=1)


# ------------------------------------------------------------------ content layouts

def _stack(ctx: Ctx, body: Box, items: list[tuple[Box, list[P]]], size: float,
           extra_max: int) -> list[tuple[Box, int]]:
    """Rows sized to their text at `size`, stacked from the top and never past the body's bottom.

    Spare height is shared out (up to extra_max per row). When the text cannot fit, the padding goes
    first, then each row gets its share of the body; its text box is then smaller than the text, so
    the text box reports the overflow. Returns (row box, text box height) pairs.
    """
    needs = [pt(fit_box(paras, box, ctx.frame, max_size=size, min_size=size).height_pt) + ctx.s(0.04)
             for box, paras in items]
    n, total = len(needs), sum(needs)
    pad = max(0, min(ctx.s(0.14), (body.h - total) // (2 * n)))
    if total <= body.h:
        extra = min(extra_max, (body.h - total - 2 * pad * n) // n)
        heights, texts = [need + 2 * pad + extra for need in needs], needs
    else:
        heights = texts = [body.h * need // total for need in needs]
    rows, y = [], body.y
    for h, text_h in zip(heights, texts):
        rows.append((Box(body.x, y, body.w, h), text_h))
        y += h
    return rows


def _numbered_rows(ctx: Ctx, slide, body: Box, grid, items: list[tuple[Box, list[P]]], size: float,
                   first_rule: str | None = None, extra_max: float = 0.45):
    """Number in column 1, text from column 2, hairline between rows. Returns (row boxes, text tops).

    extra_max is in inches: the most spare height each row may take.
    """
    rows, tops = [], []
    for i, ((row, text_h), (box, paras)) in enumerate(zip(_stack(ctx, body, items, size, ctx.s(extra_max)), items)):
        if i or first_rule:
            hline(slide, body.x, row.y, body.w, color=first_rule if i == 0 and first_rule else theme.GREY_3)
        y = row.y + (row.h - text_h) // 2
        text(ctx, slide, Box(box.x, y, box.w, text_h), paras, size=size)
        text(ctx, slide, Box(grid.x(1), y, grid.width(1), text_h), [P(str(i + 1), bold=True)],
             role="number", size=paras[0].size_at(size), color=ctx.accent, max_lines=1)
        rows.append(row)
        tops.append(y)
    return rows, tops


def executive_summary(ctx: Ctx, slide, s: dict, body: Box) -> None:
    grid = ctx.grid(body)
    rows = body.rows(len(s["points"]))
    items = []
    for row, item in zip(rows, s["points"]):
        item = {"headline": item} if isinstance(item, str) else item
        items.append((Box(grid.x(2), row.y, grid.width(11), row.h - ctx.s(0.28)),
                      _head_detail(item["headline"], item.get("detail"), 0.78)))
    size = common_size(ctx, items, max_size=18, min_size=12)
    _numbered_rows(ctx, slide, body, grid, items, size)


def two_column(ctx: Ctx, slide, s: dict, body: Box) -> None:
    grid = ctx.grid(body)
    head_h = ctx.s(0.65)
    cols = [(grid.span(1, 6), s["left"]), (grid.span(7, 6), s["right"])]
    items = [(b.inset(top=head_h + ctx.s(0.25)), [P(t, bullet=True, space_before=0.5) for t in c["bullets"]])
             for b, c in cols]
    size = common_size(ctx, items, max_size=16, min_size=12)
    for (b, c), (bullet_box, paras) in zip(cols, items):
        accent = c.get("highlight", False)  # the preferred option gets the accent heading and rule
        text(ctx, slide, b.moved(h=head_h), [P(c["heading"], bold=True)], max_size=18, min_size=14, max_lines=2,
             color=ctx.accent if accent else theme.INK, anchor="b")
        hline(slide, b.x, b.y + head_h + ctx.s(0.08), b.w, color=ctx.accent if accent else theme.GREY_3,
              width=1.5 if accent else 0.75)
        text(ctx, slide, bullet_box, paras, size=size)


def pillars(ctx: Ctx, slide, s: dict, body: Box) -> None:
    grid = ctx.grid(body)
    n = len(s["pillars"])
    boxes = grid.equal(n)
    numbered = s.get("numbered", True)
    num_h = ctx.s(0.55) if numbered else 0
    head_h = ctx.s(0.7)
    texts = []
    for b, p in zip(boxes, s["pillars"]):
        paras = ([P(p["body"])] if p.get("body") else []) + [P(t, bullet=True, space_before=0.4) for t in p.get("bullets", [])]
        texts.append((b.inset(top=num_h + head_h + ctx.s(0.25)), paras or [P("")]))
    size = common_size(ctx, texts, max_size=16, min_size=12)
    head_size = common_size(ctx, [(Box(0, 0, b.w, head_h), [P(p["heading"], bold=True)]) for b, p in zip(boxes, s["pillars"])],
                            max_size=18, min_size=14, max_lines=2)
    for i, (b, p, (tb, paras)) in enumerate(zip(boxes, s["pillars"], texts)):
        y = b.y
        if numbered:
            text(ctx, slide, Box(b.x, y, b.w, num_h), [P(f"{i + 1:02d}", bold=True)], role="number", size=20,
                 color=ctx.accent, max_lines=1, anchor="b")
            y += num_h
        text(ctx, slide, Box(b.x, y, b.w, head_h), [P(p["heading"], bold=True)], size=head_size, max_lines=2, anchor="b")
        hline(slide, b.x, y + head_h + ctx.s(0.08), b.w, color=theme.GREY_3)
        text(ctx, slide, tb, paras, size=size)


def timeline(ctx: Ctx, slide, s: dict, body: Box) -> None:
    grid = ctx.grid(body)
    ms = s["milestones"]
    cols = grid.equal(len(ms))
    when_h = ctx.s(0.45)
    line_y = body.y + when_h + ctx.s(0.2)
    dot = ctx.s(0.18)
    hline(slide, body.x, line_y, body.w, color=theme.GREY_2, width=1.5)
    below = [(Box(c.x, line_y + ctx.s(0.35), c.w, body.bottom - line_y - ctx.s(0.35)),
              _head_detail(m["label"], m.get("detail"), 0.875))
             for c, m in zip(cols, ms)]
    size = common_size(ctx, below, max_size=16, min_size=12)
    when_size = common_size(ctx, [(Box(0, 0, c.w, when_h), [P(m["when"], bold=True)]) for c, m in zip(cols, ms)],
                            max_size=16, min_size=12, max_lines=2)
    for c, m, (box, paras) in zip(cols, ms, below):
        hi = m.get("highlight", False)
        text(ctx, slide, Box(c.x, body.y, c.w, when_h), [P(m["when"], bold=True)], size=when_size,
             color=ctx.accent if hi else theme.INK_2, anchor="b", max_lines=2)
        rect(slide, Box(c.x, line_y - dot // 2, dot, dot), fill=ctx.accent if hi else theme.INK_2,
             shape=MSO_SHAPE.OVAL, name="sw:dot")
        text(ctx, slide, box, paras, size=size)


def quote(ctx: Ctx, slide, s: dict, body: Box) -> None:
    grid = ctx.grid(body)
    area = grid.span(2, 10)
    bar_w = ctx.s(0.06)
    inner = area.inset(left=bar_w + ctx.s(0.3))
    q = [P(s["quote"], head=True)]
    attribution = s.get("attribution")
    attr_h = ctx.s(0.5) if attribution else 0
    fit = fit_box(q, inner.moved(h=inner.h - attr_h), ctx.frame, max_size=28, min_size=20)
    q_h = min(inner.h - attr_h, pt(fit.height_pt) + ctx.s(0.1))
    block_h = q_h + attr_h
    top = body.y + max(0, (body.h - block_h) // 2)
    rect(slide, Box(area.x, top, bar_w, block_h), fill=ctx.accent, name="sw:accent-bar")
    text(ctx, slide, Box(inner.x, top, inner.w, q_h), q, role="quote", size=fit.size)
    if attribution:
        text(ctx, slide, Box(inner.x, top + q_h, inner.w, attr_h), [P(attribution)], size=14,
             color=theme.INK_2, anchor="m")


def recommendations(ctx: Ctx, slide, s: dict, body: Box) -> None:
    grid = ctx.grid(body)
    items = s["items"]
    extra = any(i.get("owner") or i.get("due") for i in items)
    lab = ctx.labels
    head_h = ctx.s(0.35) if extra else 0
    if extra:
        for col, label in ((9, lab["owner"]), (11, lab["due"])):
            text(ctx, slide, Box(grid.x(col), body.y, grid.width(2), head_h), [P(label, bold=True)], role="label",
                 size=12, color=theme.INK_2, anchor="b")
    rows_area = body.inset(top=head_h + (ctx.s(0.06) if extra else 0))
    span = 7 if extra else 11
    main = []
    for row, it in zip(rows_area.rows(len(items)), items):
        main.append((Box(grid.x(2), row.y, grid.width(span), row.h - ctx.s(0.28)),
                     _head_detail(it["action"], it.get("detail"), 0.875)))
    size = common_size(ctx, main, max_size=16, min_size=12)
    rows, tops = _numbered_rows(ctx, slide, rows_area, grid, main, size,
                                first_rule=theme.INK if extra else None, extra_max=0.3)
    if extra:
        for row, top, it in zip(rows, tops, items):
            for col, key in ((9, "owner"), (11, "due")):
                if it.get(key):
                    text(ctx, slide, Box(grid.x(col), top, grid.width(2), row.bottom - top), [P(it[key])],
                         size=min(size, 14))


def exhibit(ctx: Ctx, slide, s: dict, body: Box) -> None:
    exhibits.render(ctx, slide, s["exhibit"], body)


def exhibit_takeaways(ctx: Ctx, slide, s: dict, body: Box) -> None:
    grid = ctx.grid(body)
    exhibit_box = exhibits.render(ctx, slide, s["exhibit"], grid.span(1, 8))
    panel = grid.span(9, 4)
    heading = s.get("takeaways_heading") or ctx.labels["so_what"]
    head_box, rest = panel.take_top(pt(14 * 1.3))  # same top and size as the exhibit title beside it
    text(ctx, slide, head_box, [P(heading, bold=True)], size=14, color=ctx.accent, max_lines=1, role="label")
    rule_y = rest.y + ctx.s(0.06)
    hline(slide, panel.x, rule_y, panel.w, color=ctx.accent, width=2)
    # the takeaways start on the same line as the chart or table beside them
    top = max(exhibit_box.y, rule_y + ctx.s(0.08))
    text(ctx, slide, Box(panel.x, top, panel.w, panel.bottom - top),
         [P(t, bullet=True, space_before=0.6) for t in s["takeaways"]], max_size=16, min_size=12)


def table(ctx: Ctx, slide, s: dict, body: Box) -> None:
    exhibits.table(ctx, slide, s["table"], body)


def appendix(ctx: Ctx, slide, s: dict, body: Box) -> None:
    if "exhibit" in s:
        exhibits.render(ctx, slide, s["exhibit"], body)
    elif "table" in s:
        exhibits.table(ctx, slide, s["table"], body)
    else:
        text(ctx, slide, body, [P(b, bullet=True, space_before=0.5) for b in s["bullets"]], max_size=16, min_size=12)


def ghost(ctx: Ctx, slide, s: dict, body: Box) -> None:
    hints = s.get("ghost") or {}
    lab = ctx.labels
    box = body.inset(top=ctx.s(0.1))
    rect(slide, box, line=theme.GREY_2, line_w=1, dash=True, name="sw:ghost-box")
    paras = []
    if hints.get("exhibit"):
        paras.append(P(f"{lab['planned']}{lab['sep']}{hints['exhibit']}", bold=True))
    if hints.get("layout"):
        paras.append(P(f"{lab['layout']}{lab['sep']}{hints['layout']}", space_before=0.3))
    need = [e for e in hints.get("evidence", []) if DATA_NEEDED.search(e)]
    have = [e for e in hints.get("evidence", []) if not DATA_NEEDED.search(e)]
    for heading, items in ((lab["evidence"], have), (lab["still_needed"], need)):
        if items:
            paras.append(P(heading, bold=True, space_before=1.0 if paras else 0))
            paras += [P(e, bullet=True, space_before=0.3) for e in items]
    if paras:
        inner = box.inset(ctx.s(0.3), ctx.s(0.25), ctx.s(0.3), ctx.s(0.25))
        text(ctx, slide, inner, paras, max_size=16, min_size=12, color=theme.INK_2)


CONTENT = {
    "executive_summary": executive_summary, "exhibit": exhibit, "exhibit_takeaways": exhibit_takeaways,
    "two_column": two_column, "pillars": pillars, "table": table, "timeline": timeline, "quote": quote,
    "recommendations": recommendations, "appendix": appendix, "ghost": ghost,
}
FULL = {"cover": cover, "section": section}
