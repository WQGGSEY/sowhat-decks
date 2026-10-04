"""Shared drawing primitives and per-slide chrome (title, tracker, sticker, footer, page number)."""
from __future__ import annotations

from dataclasses import dataclass

from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Pt

from . import text_fit, theme
from .grid import Box, Grid, inch, pt
from .template_map import Frame, Template, box_of, insets_of, page_placeholder
from .text import FLOORS, P, Report, fit_box, write


@dataclass
class Ctx:
    """Everything a layout needs to draw one slide."""
    tpl: Template
    meta: dict
    report: Report
    slide_no: int = 0

    @property
    def frame(self) -> Frame:
        return self.tpl.frame

    @property
    def lang(self) -> str:
        return self.meta.get("language", "en")

    @property
    def labels(self) -> dict:
        return theme.LABELS[self.lang]

    @property
    def accent(self) -> str:
        return self.tpl.accent

    @property
    def accent_tint(self) -> str:
        """Light accent fill for highlighted table cells and quadrants."""
        return theme.tint(self.tpl.accent, 0.86)

    def s(self, inches: float) -> int:
        """A gap in inches, scaled to the template's slide height."""
        return int(inch(inches) * self.frame.scale)

    def grid(self, body: Box) -> Grid:
        return Grid(body, gutter=int(inch(0.2) * body.w / theme.BODY_BOX.w) if body.w else inch(0.2))


# ---------------------------------------------------------------- primitives

def text(ctx: Ctx, slide, box: Box, paras: list[P], *, role: str = "body", max_size: float = 16,
         min_size: float | None = None, align: str = "l", anchor: str = "t", max_lines: int | None = None,
         color: str | None = theme.INK, size: float | None = None):
    """Add a text box, shrink its text to fit, record the result. Returns (shape, fit).

    The role's floor (text.FLOORS) is the lowest size the text may shrink to.
    """
    floor = FLOORS[role]
    if size is not None and size < floor:
        raise ValueError(f"{role} text at {size}pt is below its {floor}pt floor")
    lo, hi = (size, size) if size is not None else (max(floor, min_size or floor), max_size)
    fit = fit_box(paras, box, ctx.frame, max_size=hi, min_size=lo, max_lines=max_lines)
    shape = slide.shapes.add_textbox(box.x, box.y, box.w, box.h)
    shape.name = f"sw:{role}"
    write(shape.text_frame, paras, fit.sizes, align=align, anchor=anchor, lang=ctx.lang, default_color=color)
    ctx.report.add(ctx.slide_no, role, fit, " / ".join(p.text for p in paras))
    return shape, fit


def common_size(ctx: Ctx, items: list[tuple[Box, list[P]]], *, max_size: float, min_size: float,
                max_lines: int | None = None) -> float:
    """Largest base size at which every (box, paras) item fits; keeps like elements equal."""
    size = max_size
    for box, paras in items:  # each item can only lower the size the earlier ones allowed
        size = fit_box(paras, box, ctx.frame, max_size=size, min_size=min_size, max_lines=max_lines).size
    return size


def rect(slide, box: Box, *, fill: str | None = None, line: str | None = None, line_w: float = 0.75,
         dash: bool = False, shape: MSO_SHAPE = MSO_SHAPE.RECTANGLE, name: str = "sw:shape"):
    sp = slide.shapes.add_shape(shape, box.x, box.y, box.w, box.h)
    sp.name = name
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = RGBColor.from_string(line)
        sp.line.width = Pt(line_w)
        if dash:
            sp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    else:
        sp.line.fill.background()
    if sp.has_text_frame:
        sp.text_frame.text = ""
    return sp


def connector(slide, x1: int, y1: int, x2: int, y2: int, *, color: str = theme.GREY_3, width: float = 0.75,
         name: str = "sw:rule", arrow: bool = False):
    """Straight connector; arrow=True puts an arrowhead at the end point."""
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    ln.name = name
    ln.line.color.rgb = RGBColor.from_string(color)
    ln.line.width = Pt(width)
    if arrow:
        ln.line._get_or_add_ln().append(parse_xml(f'<a:tailEnd {nsdecls("a")} type="triangle" w="med" len="med"/>'))
    return ln


def hline(slide, x: int, y: int, w: int, **kw):
    return connector(slide, x, y, x + w, y, **kw)


# ---------------------------------------------------------------- slide chrome

def keep_placeholders(slide, keep_types: tuple) -> None:
    """Remove placeholders the layout added that this slide does not use."""
    for ph in list(slide.placeholders):
        if ph.placeholder_format.type not in keep_types:
            ph._element.getparent().remove(ph._element)


def set_title(ctx: Ctx, slide, text_value: str, *, role: str = "title", max_size: float = 28,
              min_size: float = 24, max_lines: int = 2, box: Box | None = None):
    """Put the action title in the title placeholder (moved to box if given), fitted to it."""
    ph = slide.shapes.title
    paras = [P(text_value, head=True)]
    if ph is None:
        return text(ctx, slide, box or ctx.frame.title, paras, role=role, max_size=max_size, min_size=min_size,
                    max_lines=max_lines)
    if box is not None:
        ph.left, ph.top, ph.width, ph.height = box.x, box.y, box.w, box.h
    return ph, fill_placeholder(ctx, ph, paras, role=role, max_size=max_size, min_size=min_size, max_lines=max_lines)


def fill_placeholder(ctx: Ctx, ph, paras: list[P], *, role: str, max_size: float, min_size: float,
                     max_lines: int | None = None):
    """Fit text into a template placeholder, keeping its alignment, anchor, insets and colour."""
    fit = fit_box(paras, box_of(ph), ctx.frame, max_size=max_size, min_size=min_size, max_lines=max_lines,
                  insets=insets_of(ph))
    ph.name = f"sw:{role}"
    write(ph.text_frame, paras, fit.sizes, align=None, anchor=None, insets=None, lang=ctx.lang, default_color=None)
    ctx.report.add(ctx.slide_no, role, fit, " / ".join(p.text for p in paras))
    return fit


def tracker(ctx: Ctx, slide, label: str) -> None:
    f = ctx.frame
    h = ctx.s(0.25)
    y = max(ctx.s(0.05), f.title.y - h - ctx.s(0.05))
    text(ctx, slide, Box(f.title.x, y, int(f.title.w * 0.6), h), [P(label, bold=True)], role="tracker",
         size=11, color=ctx.accent, anchor="b")


def sticker(ctx: Ctx, slide) -> Box:
    """DRAFT sticker at the top right. Returns its box."""
    f = ctx.frame
    paras = [P(ctx.labels["draft"], bold=True)]
    pad = ctx.s(0.08)
    w = max(ctx.s(0.8), pt(text_fit.text_width(paras[0].text, f.body_font, 10, bold=True) * text_fit.SAFETY) + 2 * pad)
    h = ctx.s(0.28)
    box = Box(f.title.right - w, max(ctx.s(0.1), f.title.y - h - ctx.s(0.08)), w, h)
    sp = rect(slide, box, line=theme.INK_2, line_w=1, name="sw:sticker")
    insets = (pad, 0, pad, 0)
    fit = fit_box(paras, box, f, max_size=10, min_size=10, max_lines=1, insets=insets)
    write(sp.text_frame, paras, fit.sizes, align="c", anchor="m", insets=insets, lang=ctx.lang, default_color=theme.INK_2)
    ctx.report.add(ctx.slide_no, "sticker", fit, paras[0].text)
    return box


def page_number(ctx: Ctx, slide, layout) -> None:
    """Slide number field, in a clone of the template's slide-number placeholder when it has one."""
    layout_ph = page_placeholder(layout)
    color = ""
    if layout_ph is not None:
        slide.shapes.clone_placeholder(layout_ph)
        shape = slide.placeholders[layout_ph.placeholder_format.idx]
    else:
        f = ctx.frame
        shape = slide.shapes.add_textbox(f.page.x, f.page.y, f.page.w, f.page.h)
        write(shape.text_frame, [P("")], [10], align="r", anchor="b", lang=ctx.lang)
        color = f'<a:solidFill><a:srgbClr val="{theme.INK_2}"/></a:solidFill>'
    shape.name = "sw:page"
    shape.text_frame.paragraphs[0].clear()._p.append(parse_xml(
        f'<a:fld {nsdecls("a")} id="{{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}}" type="slidenum">'
        f'<a:rPr lang="en-US" sz="1000">{color}</a:rPr><a:t>{ctx.slide_no}</a:t></a:fld>'))


def _split_label(text: str, words: tuple[str, ...]) -> tuple[str | None, str]:
    """('Source', 'rest') if text starts with one of the words followed by ':' or '：', else (None, text)."""
    for word in words:
        head = text[:len(word)]
        tail = text[len(word):].lstrip()
        if head.lower() == word.lower() and tail[:1] in (":", "："):
            return head, tail[1:].strip()
    return None, text


def source_text(ctx: Ctx, s: dict) -> str:
    """The slide's source line.

    Explicit rules, no guessing from words inside the source:
    - An existing "Source:" / "Sources:" / "출처:" / "出所：" prefix is kept as written and never doubled.
    - When the slide or deck is marked sample_data, "Sample data" is the first item after the prefix,
      unless the source already starts with exactly that label.
    """
    lab = ctx.labels
    prefix, rest = _split_label((s.get("source") or "").strip(), ("Sources", "Source", lab["source"]))
    sample = lab["sample"]
    starts_with_sample = rest[:len(sample)].lower() == sample.lower() and rest[len(sample):][:1] in ("", ";", ".", ",")
    if s.get("sample_data", ctx.meta.get("sample_data", False)) and not starts_with_sample:
        rest = f"{sample}; {rest}" if rest else sample
    return f"{prefix or lab['source']}{lab['sep']}{rest}" if rest else ""


def footer(ctx: Ctx, slide, s: dict) -> int:
    """Footnotes and source line above the bottom edge. Returns the top y of the footer."""
    f = ctx.frame
    lines = [n for n in s.get("footnotes", []) if n]
    src = source_text(ctx, s)
    if src:
        lines.append(src)
    right = f.body.right
    if f.page.y < f.footer_bottom and f.page.bottom > f.footer_bottom - ctx.s(0.4) and f.page.x > f.slide.w / 2:
        right = f.page.x - ctx.s(0.15)
    conf = ctx.meta.get("confidentiality")
    if conf:
        cw = pt(text_fit.text_width(conf, f.body_font, 10) * 1.1) + ctx.s(0.1)
        cbox = Box(right - cw, f.footer_bottom - ctx.s(0.2), cw, ctx.s(0.2))
        text(ctx, slide, cbox, [P(conf)], role="footer", size=10, color=theme.INK_2, align="r", anchor="b")
        right = cbox.x - ctx.s(0.15)
    one_line = pt(10 * text_fit.LINE_HEIGHT_CJK)
    if not lines:
        return f.footer_bottom - one_line
    paras = [P(t) for t in lines]
    width = right - f.body.x
    measured = fit_box(paras, Box(0, 0, width, inch(10)), f, max_size=10, min_size=10)
    height = pt(measured.height_pt) + pt(2)
    box = Box(f.body.x, f.footer_bottom - height, width, height)
    text(ctx, slide, box, paras, role="source", size=10, color=theme.INK_2, anchor="b")
    return box.y


def chrome(ctx: Ctx, slide, s: dict, *, kind: str, layout) -> Box | None:
    """Sticker, tracker, title, footer and page number. Returns the body box for content slides."""
    f = ctx.frame
    draft = s.get("draft", ctx.meta.get("draft", False)) or kind == "ghost"
    title_box = None
    if draft:
        st = sticker(ctx, slide)
        if kind not in ("cover", "section") and st.bottom > f.title.y:
            title_box = Box(f.title.x, f.title.y, st.x - ctx.s(0.15) - f.title.x, f.title.h)
    if kind != "cover":
        page_number(ctx, slide, layout)
    if kind in ("cover", "section"):
        return None
    label = s.get("tracker") or (ctx.labels["appendix"] if kind == "appendix" else None)
    if label:
        tracker(ctx, slide, label)
    set_title(ctx, slide, s["title"], box=title_box)
    top = footer(ctx, slide, s)
    gap = ctx.s(0.15)
    return Box(f.body.x, f.body.y, f.body.w, max(0, top - gap - f.body.y))


def finish(slide) -> None:
    """Last pass over a built slide, for rules every shape and chart must follow.

    - Drop theme style references from drawn shapes, so nothing inherits the theme's shadow.
    - Write invertIfNegative=0 on every bar series and data point: LibreOffice draws negative
      bars as positive when it is missing.
    """
    tree = slide.shapes._spTree
    for el in tree.iter(qn("p:sp"), qn("p:cxnSp")):
        style = el.find(qn("p:style"))
        if style is not None and el.find(".//" + qn("p:ph")) is None:
            el.remove(style)
    for frame in slide.shapes:
        if frame.has_chart:
            for bar_chart in frame.chart._chartSpace.iter(qn("c:barChart")):
                for ser in bar_chart.iter(qn("c:ser")):
                    ser.get_or_add_invertIfNegative().val = False
                    for dpt in ser.iter(qn("c:dPt")):
                        if dpt.find(qn("c:invertIfNegative")) is None:
                            dpt.find(qn("c:idx")).addnext(parse_xml(f'<c:invertIfNegative {nsdecls("c")} val="0"/>'))


def notes(slide, text_value: str | None) -> None:
    if text_value:
        slide.notes_slide.notes_text_frame.text = text_value
