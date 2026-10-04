"""Open the default template or a user .pptx/.potx and map it onto deck roles.

Roles: cover (title slide), section (section header), content (title only),
blank. Geometry comes from the template's own placeholders, so content lands
inside the template's title and footer zones.
"""
from __future__ import annotations

import io
import pathlib
import zipfile
from dataclasses import dataclass

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn

from . import theme
from .grid import Box, inch

_LAYOUT_TYPES = {"cover": ("title",), "section": ("secHead",), "content": ("titleOnly",), "blank": ("blank",)}
_LAYOUT_NAMES = {"cover": ("title slide",), "section": ("section",), "content": ("title only",), "blank": ("blank",)}
TITLE_TYPES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
_PRES_MAIN = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
_TMPL_MAIN = "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"


class TemplateError(ValueError):
    pass


@dataclass
class Frame:
    """Slide geometry for one template (EMU)."""
    slide: Box
    title: Box
    body: Box             # from under the title down to the footer line
    footer_bottom: int    # bottom edge of source line and page number
    page: Box             # page number box
    head_font: str
    body_font: str
    default: bool

    @property
    def scale(self) -> float:
        """Slide height relative to the default 7.5 in (for gaps and offsets)."""
        return self.slide.h / theme.SLIDE_H


@dataclass
class Template:
    prs: object
    layouts: dict
    frame: Frame
    accent: str


def _open_package(path: pathlib.Path):
    data = path.read_bytes()
    if path.suffix.lower() == ".potx":
        src, dst = zipfile.ZipFile(io.BytesIO(data)), io.BytesIO()
        with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as out:
            for item in src.infolist():
                blob = src.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    blob = blob.replace(_TMPL_MAIN.encode(), _PRES_MAIN.encode())
                out.writestr(item, blob)
        data = dst.getvalue()
    try:
        return Presentation(io.BytesIO(data))
    except Exception as exc:  # python-pptx raises several types for bad packages
        raise TemplateError(f"{path}: not a PowerPoint file ({exc})") from exc


def _drop_slides(prs) -> None:
    """A template may ship sample slides; start from its layouts only."""
    ids = prs.slides._sldIdLst
    for sld in list(ids):
        prs.part.drop_rel(sld.get(qn("r:id")))
        ids.remove(sld)


def _find_layout(prs, role: str, wanted: str | None):
    layouts = list(prs.slide_layouts)
    if wanted:
        for layout in layouts:
            if layout.name.strip().lower() == wanted.strip().lower():
                return layout
        raise TemplateError(f"template has no layout named {wanted!r} (has: {', '.join(l.name for l in layouts)})")
    for layout in layouts:
        if layout._element.get("type") in _LAYOUT_TYPES[role]:
            return layout
    for layout in layouts:
        if any(n in layout.name.lower() for n in _LAYOUT_NAMES[role]):
            return layout
    return None


def placeholders_by_type(container) -> dict:
    """{PP_PLACEHOLDER type: placeholder} for a slide or layout (first of each type)."""
    found = {}
    for ph in container.placeholders:
        found.setdefault(ph.placeholder_format.type, ph)
    return found


def title_placeholder(container):
    phs = placeholders_by_type(container)
    return next((phs[t] for t in TITLE_TYPES if t in phs), None)


def page_placeholder(container):
    return placeholders_by_type(container).get(PP_PLACEHOLDER.SLIDE_NUMBER)


def _has_title(layout) -> bool:
    return title_placeholder(layout) is not None


def insets_of(ph) -> tuple[int, int, int, int]:
    defaults = {"lIns": 91440, "tIns": 45720, "rIns": 91440, "bIns": 45720}
    found = dict(defaults)
    chain, node = [], ph
    while node is not None and len(chain) < 4:  # slide -> layout -> master
        chain.append(node._element)
        try:
            node = node._base_placeholder
        except (AttributeError, KeyError):
            node = None
    for el in reversed(chain):  # master first, then layout, then slide overrides
        txbody = el.find(qn("p:txBody"))
        body = txbody.find(qn("a:bodyPr")) if txbody is not None else None
        if body is not None:
            for k in defaults:
                if body.get(k) is not None:
                    found[k] = int(body.get(k))
    return found["lIns"], found["tIns"], found["rIns"], found["bIns"]


def box_of(shape) -> Box | None:
    """A shape's box in EMU (inherited from its layout for placeholders), or None if it has none."""
    try:
        if shape is None or shape.left is None or shape.width is None:
            return None
        return Box(int(shape.left), int(shape.top), int(shape.width), int(shape.height))
    except (AttributeError, KeyError):
        return None


def open_template(path: str | pathlib.Path | None = None, *, accent: str | None = None,
                  language: str = "en", layout_map: dict | None = None) -> Template:
    """Load the template and work out layouts and geometry."""
    if path is None:
        prs = theme.make_default_template(accent or theme.ACCENT, language)
        default = True
    else:
        path = pathlib.Path(path)
        if not path.is_file():
            raise TemplateError(f"template not found: {path}")
        prs = _open_package(path)
        _drop_slides(prs)
        theme.set_east_asian_font(prs, language)
        default = False

    layout_map = layout_map or {}
    layouts = {role: _find_layout(prs, role, layout_map.get(role)) for role in ("cover", "section", "content", "blank")}
    if layouts["content"] is None or not _has_title(layouts["content"]):
        layouts["content"] = next((l for l in prs.slide_layouts if _has_title(l)), None)
    if layouts["content"] is None:
        raise TemplateError("template has no layout with a title placeholder")
    for role in ("cover", "section"):
        if layouts[role] is None or not _has_title(layouts[role]):
            layouts[role] = layouts["content"]

    sw, sh = int(prs.slide_width), int(prs.slide_height)
    slide = Box(0, 0, sw, sh)
    title_ph = title_placeholder(layouts["content"])
    title = box_of(title_ph) or Box(int(sw * 0.04), int(sh * 0.07), int(sw * 0.92), int(sh * 0.14))
    scale = sh / theme.SLIDE_H
    page = box_of(page_placeholder(layouts["content"]))
    if page is not None and page.bottom > sh * 0.8:
        footer_bottom = page.bottom
    else:
        footer_bottom = sh - int(inch(0.3) * scale)
        if page is None:
            page = Box(title.right - int(inch(1.0) * scale), footer_bottom - int(inch(0.25) * scale),
                       int(inch(1.0) * scale), int(inch(0.25) * scale))
    body_top = title.bottom + int(inch(0.2) * scale)
    body = Box(title.x, body_top, title.w, max(0, footer_bottom - body_top))

    fonts = theme.theme_fonts(prs)
    frame = Frame(
        slide=slide, title=title,
        body=body, footer_bottom=footer_bottom, page=page,
        head_font=fonts["major"] or theme.HEAD_FONT, body_font=fonts["minor"] or theme.BODY_FONT,
        default=default,
    )
    return Template(prs=prs, layouts=layouts, frame=frame, accent=theme.theme_accent(prs).upper())
