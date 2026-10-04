"""Default design: palette, fonts, labels, and the built-in 16:9 template.

Design decisions (see references/design-system.md):
- One accent (cobalt) for the single thing that matters on a slide; everything
  else is ink or grey.
- Two fonts that ship with Windows, macOS and PowerPoint: Georgia for action
  titles, Arial for everything else. East Asian theme font per language.
"""
from __future__ import annotations

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

from .grid import Box, Grid, inch

ACCENT = "1F5AA6"
INK = "1E2329"         # primary text
INK_2 = "5A6472"       # secondary text, source lines
GREY_1 = "8C95A1"      # secondary data series
GREY_2 = "B7BEC7"      # muted data (non-highlighted bars)
GREY_3 = "D5DAE0"      # rules and dividers
GREY_4 = "F2F4F6"      # light fills
WHITE = "FFFFFF"
GREY_RAMP = [INK_2, "737D8A", GREY_1, GREY_2, GREY_3]  # series that are not highlighted, dark to light

HEAD_FONT = "Georgia"
BODY_FONT = "Arial"
EA_FONT = {"en": "", "ko": "Malgun Gothic", "ja": "Yu Gothic"}
EA_SCRIPT = {"ko": "Hang", "ja": "Jpan"}  # theme <a:font script=...> for each language
LANG_TAG = {"en": "en-US", "ko": "ko-KR", "ja": "ja-JP"}

LABELS = {
    "en": {"source": "Source", "note": "Note", "draft": "DRAFT", "sample": "Sample data",
           "appendix": "Appendix", "so_what": "So what", "owner": "Owner", "due": "Timing",
           "planned": "Planned exhibit", "evidence": "Evidence", "still_needed": "Still needed", "layout": "Layout",
           "sep": ": ", "total": "Total"},
    "ko": {"source": "출처", "note": "주", "draft": "초안", "sample": "샘플 데이터",
           "appendix": "부록", "so_what": "시사점", "owner": "담당", "due": "시기",
           "planned": "계획한 엑시빗", "evidence": "확보한 근거", "still_needed": "아직 필요한 근거", "layout": "레이아웃",
           "sep": ": ", "total": "합계"},
    "ja": {"source": "出所", "note": "注", "draft": "ドラフト", "sample": "サンプルデータ",
           "appendix": "付録", "so_what": "示唆", "owner": "担当", "due": "時期",
           "planned": "予定の図表", "evidence": "確認済みの根拠", "still_needed": "まだ必要な根拠", "layout": "レイアウト",
           "sep": "：", "total": "合計"},
}

def tint(hex_color: str, amount: float) -> str:
    """Mix a color with white. amount=0.86 keeps 14% of the color."""
    r, g, b = (int(hex_color[i:i + 2], 16) for i in (0, 2, 4))
    return "".join(f"{round(c + (255 - c) * amount):02X}" for c in (r, g, b))


SLIDE_W, SLIDE_H = 12192000, 6858000  # 13.333 x 7.5 in (16:9)
MARGIN = inch(0.5)
TITLE_BOX = Box(MARGIN, inch(0.55), SLIDE_W - 2 * MARGIN, inch(1.0))
BODY_BOX = Box(MARGIN, inch(1.75), SLIDE_W - 2 * MARGIN, inch(5.1))
PAGE_BOX = Box(SLIDE_W - MARGIN - inch(1.0), inch(6.95), inch(1.0), inch(0.25))

def _theme_part(prs):
    return prs.slide_master.part.part_related_by(RT.THEME)


def _theme_xml(prs):
    return etree.fromstring(_theme_part(prs).blob)


def _save_theme(prs, root) -> None:
    _theme_part(prs)._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def theme_fonts(prs) -> dict:
    """Latin and East Asian typefaces of the theme: {'major','minor','major_ea','minor_ea'}."""
    root = _theme_xml(prs)
    out = {}
    for role in ("major", "minor"):
        node = root.find(".//" + qn(f"a:{role}Font"))
        latin = node.find(qn("a:latin")) if node is not None else None
        ea = node.find(qn("a:ea")) if node is not None else None
        out[role] = latin.get("typeface") if latin is not None else BODY_FONT
        out[f"{role}_ea"] = ea.get("typeface") if ea is not None else ""
    return out


def theme_accent(prs) -> str:
    root = _theme_xml(prs)
    node = root.find(".//" + qn("a:clrScheme") + "/" + qn("a:accent1"))
    if node is not None and len(node):
        return node[0].get("val") or node[0].get("lastClr") or ACCENT
    return ACCENT


def set_east_asian_font(prs, language: str, force: bool = False) -> None:
    """Set the theme's East Asian font for Korean/Japanese decks.

    A user template keeps a font it already sets unless force=True.
    """
    face = EA_FONT.get(language, "")
    if not face:
        return
    root = _theme_xml(prs)
    script = EA_SCRIPT[language]
    for role in ("majorFont", "minorFont"):
        node = root.find(".//" + qn(f"a:{role}"))
        if node is None:
            continue
        ea = node.find(qn("a:ea"))
        if ea is None:
            ea = etree.SubElement(node, qn("a:ea"))
        if force or not ea.get("typeface"):
            ea.set("typeface", face)
        for f in node.findall(qn("a:font")):
            if f.get("script") == script and (force or not f.get("typeface")):
                f.set("typeface", face)
    _save_theme(prs, root)


def _set_color_scheme(root, accent: str) -> None:
    scheme = root.find(".//" + qn("a:clrScheme"))
    scheme.set("name", "SoWhat Decks")
    colors = {"dk1": INK, "lt1": WHITE, "dk2": INK_2, "lt2": GREY_4, "accent1": accent,
              "accent2": GREY_1, "accent3": GREY_2, "accent4": INK_2, "accent5": GREY_3,
              "accent6": INK, "hlink": accent, "folHlink": INK_2}
    for name, val in colors.items():
        node = scheme.find(qn(f"a:{name}"))
        del node[:]
        etree.SubElement(node, qn("a:srgbClr"), val=val)


def _set_font_scheme(root) -> None:
    scheme = root.find(".//" + qn("a:fontScheme"))
    scheme.set("name", "SoWhat Decks")
    for role, face in (("majorFont", HEAD_FONT), ("minorFont", BODY_FONT)):
        node = scheme.find(qn(f"a:{role}"))
        node.find(qn("a:latin")).set("typeface", face)
        by_script = {script: EA_FONT[lang] for lang, script in EA_SCRIPT.items()}
        for f in node.findall(qn("a:font")):
            if f.get("script") in by_script:
                f.set("typeface", by_script[f.get("script")])


def _clear_xfrm(sp) -> None:
    sppr = sp.find(qn("p:spPr"))
    xfrm = sppr.find(qn("a:xfrm")) if sppr is not None else None
    if xfrm is not None:
        sppr.remove(xfrm)


def _body_pr(sp, anchor: str = "t") -> None:
    body = sp.find(qn("p:txBody")).find(qn("a:bodyPr"))
    for k in ("lIns", "tIns", "rIns", "bIns"):
        body.set(k, "0")
    body.set("anchor", anchor)
    body.set("wrap", "square")
    for child in list(body):
        if child.tag in (qn("a:normAutofit"), qn("a:spAutoFit"), qn("a:noAutofit")):
            body.remove(child)
    etree.SubElement(body, qn("a:noAutofit"))


def _lvl(tag: str, size: int, color: str = "tx1", algn: str = "l", bold: bool = False,
         font: str = "mn", bullet: str | None = None, mar: int = 0, indent: int = 0,
         space_before: int = 0, line: int | None = None) -> etree._Element:
    lvl = etree.Element(qn(f"a:{tag}"), algn=algn, marL=str(mar), indent=str(indent))
    if line:
        etree.SubElement(etree.SubElement(lvl, qn("a:lnSpc")), qn("a:spcPct"), val=str(line))
    etree.SubElement(etree.SubElement(lvl, qn("a:spcBef")), qn("a:spcPts"), val=str(space_before * 100))  # 1/100 pt
    if bullet:
        etree.SubElement(lvl, qn("a:buFont"), typeface=BODY_FONT)
        etree.SubElement(lvl, qn("a:buChar"), char=bullet)
    else:
        etree.SubElement(lvl, qn("a:buNone"))
    rpr = etree.SubElement(lvl, qn("a:defRPr"), sz=str(size * 100), b="1" if bold else "0")
    fill = etree.SubElement(rpr, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:schemeClr"), val=color)
    etree.SubElement(rpr, qn("a:latin"), typeface=f"+{font}-lt")
    etree.SubElement(rpr, qn("a:ea"), typeface=f"+{font}-ea")
    etree.SubElement(rpr, qn("a:cs"), typeface=f"+{font}-cs")
    return lvl


def _set_lst_style(sp, *levels: etree._Element) -> None:
    txbody = sp.find(qn("p:txBody"))
    lst = txbody.find(qn("a:lstStyle"))
    del lst[:]
    for lvl in levels:
        lst.append(lvl)


def _style_placeholder(ph, box: Box, anchor: str, level) -> None:
    """Position a master/layout placeholder, zero its insets, and set its first-level text style."""
    ph.left, ph.top, ph.width, ph.height = box.x, box.y, box.w, box.h
    _body_pr(ph._element, anchor)
    if level is not None:
        _set_lst_style(ph._element, level)


def _set_master_styles(master) -> None:
    styles = master._element.find(qn("p:txStyles"))
    title = styles.find(qn("p:titleStyle"))
    del title[:]
    title.append(_lvl("lvl1pPr", 28, font="mj", line=90000))
    body = styles.find(qn("p:bodyStyle"))
    del body[:]
    body.append(_lvl("lvl1pPr", 16, bullet="•", mar=inch(0.22), indent=-inch(0.22), space_before=6))
    body.append(_lvl("lvl2pPr", 14, bullet="–", mar=inch(0.45), indent=-inch(0.2), space_before=3))
    for n in range(3, 10):
        body.append(_lvl(f"lvl{n}pPr", 14, bullet="–", mar=inch(0.45 + 0.2 * (n - 2)), indent=-inch(0.2), space_before=3))


def make_default_template(accent: str = ACCENT, language: str = "en"):
    """Build the 16:9 default template in memory from python-pptx's base package."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H

    root = _theme_xml(prs)
    root.set("name", "SoWhat Decks")
    _set_color_scheme(root, accent.lstrip("#").upper())
    _set_font_scheme(root)
    _save_theme(prs, root)
    set_east_asian_font(prs, language, force=True)

    master = prs.slide_master
    _set_master_styles(master)
    PH = PP_PLACEHOLDER
    footer = lambda align: _lvl("lvl1pPr", 10, color="tx2", algn=align)  # noqa: E731
    master_style = {
        PH.TITLE: (TITLE_BOX, "t", None),
        PH.BODY: (BODY_BOX, "t", None),
        PH.DATE: (Box(MARGIN, PAGE_BOX.y, inch(2.0), PAGE_BOX.h), "b", footer("l")),
        PH.FOOTER: (Box(SLIDE_W // 2 - inch(2.0), PAGE_BOX.y, inch(4.0), PAGE_BOX.h), "b", footer("ctr")),
        PH.SLIDE_NUMBER: (PAGE_BOX, "b", footer("r")),
    }
    for ph in master.placeholders:
        if ph.placeholder_format.type in master_style:
            _style_placeholder(ph, *master_style[ph.placeholder_format.type])

    keep = {"Title Slide", "Title and Content", "Section Header", "Title Only", "Blank"}
    for layout in list(prs.slide_layouts):
        if layout.name not in keep:
            prs.slide_layouts.remove(layout)

    grid = Grid(BODY_BOX)
    layout_style = {
        ("Title Slide", PH.CENTER_TITLE): (Box(MARGIN, inch(2.3), grid.width(9), inch(2.0)), "b",
                                           _lvl("lvl1pPr", 40, font="mj", line=90000)),
        ("Title Slide", PH.SUBTITLE): (Box(MARGIN, inch(4.55), grid.width(9), inch(1.3)), "t",
                                       _lvl("lvl1pPr", 20, color="tx2")),
        ("Section Header", PH.TITLE): (Box(grid.x(3), inch(2.4), grid.width(10), inch(1.6)), "b",
                                       _lvl("lvl1pPr", 36, font="mj", line=90000)),
        ("Section Header", PH.BODY): (Box(grid.x(3), inch(4.2), grid.width(10), inch(1.2)), "t",
                                      _lvl("lvl1pPr", 18, color="tx2")),
    }
    for layout in prs.slide_layouts:
        for ph in layout.placeholders:
            _clear_xfrm(ph._element)  # inherit the master's position unless styled below
            spec = layout_style.get((layout.name, ph.placeholder_format.type))
            if spec:
                _style_placeholder(ph, *spec)
    return prs
