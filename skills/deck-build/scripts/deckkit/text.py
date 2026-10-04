"""Write fitted text into shapes, with theme fonts and language tags."""
from __future__ import annotations

from dataclasses import dataclass, field

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

from . import text_fit
from .theme import LANG_TAG
from .grid import Box, to_pt
from .text_fit import Para

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}
BULLET_INDENT_PT = 14

# Minimum sizes by shape role. Shapes are named "sw:<role>" so checks can find them.
FLOORS = {"title": 24, "cover-title": 28, "section-title": 28, "body": 12, "label": 12,
          "exhibit-title": 12, "quote": 20, "number": 12,
          "source": 10, "page": 10, "tracker": 10, "sticker": 10, "footer": 10}  # footnotes share "source"


@dataclass
class P(Para):
    """A paragraph to write: text_fit.Para plus how it looks."""
    color: str | None = None
    bullet: bool = False
    head: bool = False       # heading font (theme major) instead of body font
    align: str | None = None


@dataclass
class Report:
    """What the builder did and what did not fit."""
    texts: list[dict] = field(default_factory=list)

    def add(self, slide_no: int, role: str, fit: text_fit.Fit, text: str) -> None:
        self.texts.append({"slide": slide_no, "role": role, "size": fit.size, "lines": fit.lines,
                           "overflow": fit.overflow, "text": text[:60]})

    @property
    def overflows(self) -> list[dict]:
        return [t for t in self.texts if t["overflow"]]

    def as_dict(self) -> dict:
        return {"overflow_count": len(self.overflows), "overflows": self.overflows, "texts": self.texts}


def lang_tag(text: str, deck_lang: str) -> str:
    """Language of a run: Hangul is Korean, kana is Japanese, other CJK follows the deck."""
    for ch in text:
        if text_fit.is_hangul(ch):
            return LANG_TAG["ko"]
        if 0x3040 <= ord(ch) <= 0x30FF:
            return LANG_TAG["ja"]
    return LANG_TAG.get(deck_lang, "en-US") if text_fit.has_cjk(text) else "en-US"


def style_run(run, size: float, *, bold: bool = False, color: str | None = None,
              head: bool = False, lang: str = "en-US") -> None:
    font = run.font
    font.size = Pt(size)
    font.bold = bold
    if color:
        font.color.rgb = RGBColor.from_string(color)
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", lang)
    if lang != "en-US":
        rpr.set("altLang", "en-US")
    for tag in ("a:latin", "a:ea", "a:cs"):
        old = rpr.find(qn(tag))
        if old is not None:
            rpr.remove(old)
    kind = "mj" if head else "mn"
    for tag in ("latin", "ea", "cs"):
        etree.SubElement(rpr, qn(f"a:{tag}"), typeface=f"+{kind}-{tag if tag != 'latin' else 'lt'}")


def _set_bullet(paragraph, bullet: bool) -> None:
    """Pin the indent the text was measured with, so a template's inherited indent cannot change it."""
    ppr = paragraph._p.get_or_add_pPr()
    indent = int(Pt(BULLET_INDENT_PT)) if bullet else 0
    ppr.set("marL", str(indent))
    ppr.set("indent", str(-indent))
    for tag in ("a:buNone", "a:buChar", "a:buFont", "a:buAutoNum"):
        old = ppr.find(qn(tag))
        if old is not None:
            ppr.remove(old)
    if bullet:
        etree.SubElement(ppr, qn("a:buFont"), typeface="Arial")
        etree.SubElement(ppr, qn("a:buChar"), char="•")
    else:
        etree.SubElement(ppr, qn("a:buNone"))


def write(text_frame, paras: list[P], sizes: list[float], *, align: str | None = "l", anchor: str | None = "t",
          insets: tuple[int, int, int, int] | None = (0, 0, 0, 0), lang: str = "en", default_color: str | None = None) -> None:
    """Replace the text frame's content with the paragraphs at the given sizes.

    align, anchor or insets set to None keep what the shape inherits (used for template placeholders).
    """
    tf = text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    if insets is not None:
        tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = insets
    if anchor:
        tf.vertical_anchor = ANCHOR[anchor]
    tf.clear()
    for i, (para, size) in enumerate(zip(paras, sizes)):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if para.align or align:
            p.alignment = ALIGN[para.align or align]
        p.space_before = Pt(round(para.space_before * size, 1) if i else 0)  # as measured
        _set_bullet(p, para.bullet)
        lines = para.text.split("\n")
        for j, line in enumerate(lines):
            if j:
                p.add_line_break()
            run = p.add_run()
            run.text = line
            style_run(run, size, bold=para.bold, color=para.color or default_color,
                      head=para.head, lang=lang_tag(line, lang))


def fit_box(paras: list[P], box: Box, frame, *, max_size: float, min_size: float,
            max_lines: int | None = None, insets: tuple[int, int, int, int] = (0, 0, 0, 0)) -> text_fit.Fit:
    """Fit paragraphs into a box, measuring each with the template's heading or body font."""
    for p in paras:
        p.font = frame.head_font if p.head else frame.body_font
        p.indent = BULLET_INDENT_PT if p.bullet else 0
    width = to_pt(box.w - insets[0] - insets[2])
    height = to_pt(box.h - insets[1] - insets[3])
    return text_fit.fit(paras, width, height, max_size, min_size, frame.body_font, max_lines)
