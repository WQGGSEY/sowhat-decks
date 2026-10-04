"""Inspect a built deck: bounds, type-size floors, title length, estimated overflow, native charts.

Works on the saved file (re-opened with python-pptx), so it checks what the
user gets, not what the builder intended.
"""
from __future__ import annotations

from dataclasses import dataclass

from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.shapes.group import GroupShape

from . import text_fit
from .grid import Box, pt, to_pt
from .template_map import TITLE_TYPES, box_of, insets_of
from .text import FLOORS
from .theme import BODY_FONT, HEAD_FONT, theme_fonts

FOOTER_TYPES = (PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE)

@dataclass
class Issue:
    slide: int
    shape: str
    kind: str
    detail: str

    def __str__(self) -> str:
        return f"slide {self.slide} {self.shape}: {self.kind} ({self.detail})"


def _role(shape) -> str:
    name = shape.name or ""
    if name.startswith("sw:"):
        return name[3:]
    if shape.is_placeholder and shape.placeholder_format.type in TITLE_TYPES:
        return "title"
    if shape.is_placeholder and shape.placeholder_format.type in FOOTER_TYPES:
        return "page"
    return "body"


def _iter(shapes, group=None):
    """(shape, the top-level group it sits in or None), groups flattened."""
    for sh in shapes:
        if isinstance(sh, GroupShape):
            yield from _iter(sh.shapes, group or sh)
        else:
            yield sh, group


_RUN_TAGS = (qn("a:r"), qn("a:fld"))


def _runs(paragraph):
    """(text, rPr) for every text run and field (slide number) in a paragraph."""
    for el in paragraph._p:
        if el.tag in _RUN_TAGS:
            yield el.findtext(qn("a:t")) or "", el.find(qn("a:rPr"))


def _size(rpr) -> float | None:
    return int(rpr.get("sz")) / 100 if rpr is not None and rpr.get("sz") else None


def _run_sizes(text_frame) -> list[float]:
    return [_size(rpr) for p in text_frame.paragraphs for t, rpr in _runs(p) if t.strip() and _size(rpr)]


def _paras(text_frame, head_font: str, body_font: str) -> list[text_fit.Para]:
    """The frame's paragraphs as text_fit paragraphs at their written sizes."""
    out = []
    for p in text_frame.paragraphs:
        runs = list(_runs(p))
        sizes = [_size(rpr) for _, rpr in runs if _size(rpr)]
        text = p.text.replace("\v", "\n")  # python-pptx shows line breaks as vertical tabs
        if not text.strip() or not sizes:
            continue
        first = runs[0][1]
        latin = first.find(qn("a:latin")) if first is not None else None
        ppr = p._p.find(qn("a:pPr"))
        size = max(sizes)
        out.append(text_fit.Para(
            text, bold=first is not None and first.get("b") == "1", size=size,
            font=head_font if latin is not None and latin.get("typeface", "").startswith("+mj") else body_font,
            indent=to_pt(int(ppr.get("marL", 0))) if ppr is not None else 0,
            space_before=(p.space_before.pt / size) if p.space_before is not None else 0))
    return out


def estimate(shape, head_font: str, body_font: str) -> tuple[float, float, int]:
    """(needed height pt, available height pt, line count) for a text shape."""
    l, t, r, b = insets_of(shape)
    paras = _paras(shape.text_frame, head_font, body_font)
    _, counts, needed = text_fit.measure(paras, 0, to_pt(int(shape.width) - l - r), body_font)
    return needed, to_pt(int(shape.height) - t - b), sum(counts)


def _table_issues(no: int, shape, body_font: str) -> list[Issue]:
    issues, total = [], 0.0
    widths = [int(c.width) for c in shape.table.columns]
    for row in shape.table.rows:
        row_h = 0.0
        for cell, w in zip(row.cells, widths):
            paras = _paras(cell.text_frame, body_font, body_font)
            if not paras:
                continue
            if min(p.size for p in paras) < FLOORS["body"]:
                issues.append(Issue(no, shape.name, "font-too-small", f"table cell {min(p.size for p in paras)}pt"))
            width = to_pt(w - int(cell.margin_left) - int(cell.margin_right))
            height = text_fit.measure(paras, 0, width, body_font)[2]
            row_h = max(row_h, height + to_pt(int(cell.margin_top) + int(cell.margin_bottom)))
        total += row_h
    if total > to_pt(int(shape.height)) + 1:
        issues.append(Issue(no, shape.name, "overflow", f"table needs {total:.0f}pt, frame {to_pt(int(shape.height)):.0f}pt"))
    return issues


def _chart_issues(no: int, shape) -> list[Issue]:
    issues = []
    xml = shape.chart._chartSpace
    for el in xml.iter(qn("a:defRPr"), qn("a:rPr")):
        if el.get("sz") and int(el.get("sz")) < FLOORS["label"] * 100:
            issues.append(Issue(no, shape.name, "font-too-small", f"chart text {int(el.get('sz')) / 100}pt"))
    if shape.chart.has_legend:
        issues.append(Issue(no, shape.name, "legend", "use direct labels instead of a legend"))
    return issues


def _text_area(shape, box: Box, needed_pt: float) -> Box:
    """Where the text actually sits: the box cut to the text's estimated height at its anchor.

    Placeholders inherit their anchor from the template, so their whole box counts.
    """
    if shape.is_placeholder:
        return box
    _, top, _, bottom = insets_of(shape)
    h = min(box.h, pt(needed_pt) + top + bottom)
    anchor = shape.text_frame._txBody.find(qn("a:bodyPr")).get("anchor", "t")
    y = box.y if anchor == "t" else box.bottom - h if anchor == "b" else box.y + (box.h - h) // 2
    return Box(box.x, y, box.w, h)


def _overlaps(no: int, texts: list[tuple[object, Box, object]]) -> list[Issue]:
    """Text that covers other text (more than 0.02 in both ways), unless both sit in one group."""
    issues, tol = [], 18288  # 0.02 in
    for i, (a, ba, ga) in enumerate(texts):
        for b, bb, gb in texts[i + 1:]:
            if ga is not None and ga is gb:
                continue
            w = min(ba.right, bb.right) - max(ba.x, bb.x)
            h = min(ba.bottom, bb.bottom) - max(ba.y, bb.y)
            if w > tol and h > tol:
                issues.append(Issue(no, b.name, "overlap", f"covers {a.name} by {w / 914400:.2f} x {h / 914400:.2f} in"))
    return issues


def inspect(prs) -> list[Issue]:
    """Every rule violation in the deck. Empty list = clean."""
    try:
        fonts = theme_fonts(prs)
        head_font, body_font = fonts["major"] or HEAD_FONT, fonts["minor"] or BODY_FONT
    except (KeyError, ValueError, AttributeError):  # no readable theme: check with the default fonts
        head_font, body_font = HEAD_FONT, BODY_FONT
    slide_box = Box(0, 0, int(prs.slide_width), int(prs.slide_height))
    issues: list[Issue] = []
    for no, slide in enumerate(prs.slides, 1):
        texts = []
        for shape, group in _iter(slide.shapes):
            name = shape.name or shape.shape_type
            box = box_of(shape)
            if box is not None and not box.inside(slide_box, tolerance=1):
                issues.append(Issue(no, name, "out-of-bounds", f"{box}"))
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                issues.append(Issue(no, name, "picture", "exhibits must be native charts or shapes"))
            if shape.has_chart:
                issues.extend(_chart_issues(no, shape))
            if shape.has_table:
                issues.extend(_table_issues(no, shape, body_font))
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            role = _role(shape)
            floor = FLOORS.get(role, FLOORS["body"])
            sizes = _run_sizes(shape.text_frame)
            if sizes and min(sizes) < floor:
                issues.append(Issue(no, name, "font-too-small", f"{min(sizes)}pt < {floor}pt for {role}"))
            if box is None or not sizes:
                continue
            needed, avail, lines = estimate(shape, head_font, body_font)
            texts.append((shape, _text_area(shape, box, needed), group))
            if needed > avail + 0.5:
                issues.append(Issue(no, name, "overflow", f"needs {needed:.0f}pt, box {avail:.0f}pt"))
            if role == "title" and lines > 2:
                issues.append(Issue(no, name, "title-lines", f"{lines} lines"))
        issues.extend(_overlaps(no, texts))
    return issues
