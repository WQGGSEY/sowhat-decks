"""Parse any .pptx into a plain, JSON-serialisable dict.

Everything the checks need is resolved here, once: effective font sizes (through
placeholder, master and theme inheritance), font names, colors, positions of
shapes inside groups, an estimate of how much room each text needs, and simple
image statistics that tell a chart picture from a photo.

All lengths are EMU (914400 per inch, 12700 per point) unless the key ends in
_pt or _in. Each shape has `bbox` (the frame stored in the file, [x, y, w, h])
and `extent` (the area it is drawn in: taller than bbox when a grow-to-fit text
box or a table row grows to fit its text).
"""

from __future__ import annotations

import colorsys
import io
import re
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn

from . import textfit

EMU_PER_IN = 914400
EMU_PER_PT = 12700
DEFAULT_SIZE_PT = 18.0
DEFAULT_INSETS = (91440, 45720, 91440, 45720)  # left, top, right, bottom
MAX_UNZIPPED = 1 << 30  # refuse decks that unzip to more than 1 GiB (zip bombs)
MAX_IMAGE_BYTES = 40 << 20  # skip image statistics for images over 40 MB

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
}
URI_CHART = "http://schemas.openxmlformats.org/drawingml/2006/chart"
URI_TABLE = "http://schemas.openxmlformats.org/drawingml/2006/table"
URI_DIAGRAM = "http://schemas.openxmlformats.org/drawingml/2006/diagram"

SOURCE_RE = re.compile(
    r"^\s*(\(?\d\)?\s*)?(sources?|note|notes|data|footnote|source and notes|"
    r"출처|자료|주|참고|出典|出所|注|データ)\s*[:：]", re.IGNORECASE)
SOURCE_ANYWHERE_RE = re.compile(r"\b(sources?)\s*:|(출처|자료)\s*[:：]|(出典|出所)\s*[:：]|"
                                r"\bsample data\b|샘플 데이터|サンプルデータ", re.IGNORECASE)

TITLE_TYPES = {"title", "ctrTitle"}
FOOTER_TYPES = {"dt", "ftr", "sldNum", "hdr"}
PRESET_COLORS = {"black": "000000", "white": "FFFFFF", "red": "FF0000", "green": "008000",
                 "blue": "0000FF", "yellow": "FFFF00", "gray": "808080", "grey": "808080",
                 "orange": "FFA500"}


def _xp(el, expr):
    return etree._Element.xpath(el, expr, namespaces=NS)


def _local(el) -> str:
    return etree.QName(el).localname if isinstance(el.tag, str) else ""


def _int(v, default=None):
    try:
        return int(v)
    except (TypeError, ValueError):
        return default


# ------------------------------------------------------------------ theme


class Theme:
    def __init__(self, master):
        self.major = self.minor = None
        self.colors: dict[str, str] = {}
        self.clrmap = {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}
        try:
            cm = master._element.find(qn("p:clrMap"))
            if cm is not None:
                self.clrmap.update(dict(cm.attrib))
        except Exception:
            pass
        try:
            root = parse_xml(master.part.part_related_by(RT.THEME).blob)
        except Exception:
            return
        for kind in ("major", "minor"):
            latin = _xp(root, f".//a:fontScheme/a:{kind}Font/a:latin/@typeface")
            setattr(self, kind, latin[0] if latin and latin[0] else None)
        for el in _xp(root, ".//a:clrScheme/*"):
            val = None
            for c in el:
                if _local(c) == "srgbClr":
                    val = c.get("val")
                elif _local(c) == "sysClr":
                    val = c.get("lastClr")
            if val:
                self.colors[_local(el)] = val.upper()

    def font(self, typeface):
        if not typeface:
            return None
        if typeface.startswith("+mj"):
            return self.major
        if typeface.startswith("+mn"):
            return self.minor
        return typeface

    def color(self, el):
        """Resolve a color element (srgbClr, schemeClr, ...) to 'RRGGBB'."""
        tag = _local(el)
        val = None
        if tag == "srgbClr":
            val = (el.get("val") or "").upper()
        elif tag == "schemeClr":
            name = el.get("val")
            if name == "phClr":
                return None
            name = self.clrmap.get(name, name)
            val = self.colors.get(name)
        elif tag == "sysClr":
            val = el.get("lastClr")
        elif tag == "prstClr":
            val = PRESET_COLORS.get((el.get("val") or "").lower())
        elif tag == "scrgbClr":
            try:
                val = "".join(f"{min(255, int(int(el.get(k)) / 100000 * 255)):02X}"
                              for k in ("r", "g", "b"))
            except (TypeError, ValueError):
                val = None
        if not val or len(val) != 6:
            return None
        return _apply_lum(val.upper(), el)


def _apply_lum(rgb: str, el) -> str:
    mod = off = None
    for c in el:
        if _local(c) == "lumMod":
            mod = _int(c.get("val"))
        elif _local(c) == "lumOff":
            off = _int(c.get("val"))
    if mod is None and off is None:
        return rgb
    r, g, b = (int(rgb[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = min(1.0, max(0.0, l * (mod or 100000) / 100000 + (off or 0) / 100000))
    r, g, b = colorsys.hls_to_rgb(h, l, s)
    return "".join(f"{int(round(v * 255)):02X}" for v in (r, g, b))


def _solid_color(parent, theme):
    """Color of a:solidFill (or first gradient stop) directly under parent."""
    if parent is None:
        return None
    sf = parent.find(qn("a:solidFill"))
    if sf is not None and len(sf):
        return theme.color(sf[0])
    gs = _xp(parent, "./a:gradFill/a:gsLst/a:gs[1]/*[1]")
    if gs:
        return theme.color(gs[0])
    return None


# ------------------------------------------------------------------ placeholders


def _ph(el):
    found = _xp(el, "./*[1]/p:nvPr/p:ph")
    if not found:
        return None
    ph = found[0]
    return {"type": ph.get("type", "obj"), "idx": _int(ph.get("idx"), 0)}


def _find_ph(container_el, ph, by_idx=True):
    """Find the matching placeholder element in a layout or master spTree."""
    if container_el is None or ph is None:
        return None
    candidates = []
    for el in _xp(container_el, ".//p:cSld/p:spTree/*"):
        info = _ph(el)
        if info:
            candidates.append((el, info))
    if by_idx:
        for el, info in candidates:
            if info["idx"] == ph["idx"] and ph["idx"] != 0:
                return el
    want = _master_type(ph["type"])
    for el, info in candidates:
        if _master_type(info["type"]) == want:
            return el
    return None


def _master_type(t):
    if t in TITLE_TYPES:
        return "title"
    if t in FOOTER_TYPES:
        return t
    return "body"


def _txstyle(master_el, ph_type):
    if master_el is None:
        return None
    if ph_type in TITLE_TYPES:
        name = "titleStyle"
    elif ph_type in FOOTER_TYPES or ph_type is None:
        name = "otherStyle"
    else:
        name = "bodyStyle"
    found = _xp(master_el, f"./p:txStyles/p:{name}")
    return found[0] if found else None


class StyleChain:
    """List styles and bodyPr elements a text body inherits from, nearest first."""

    def __init__(self, el, slide_ctx):
        self.list_styles = []
        self.body_prs = []
        ls = _xp(el, "./p:txBody/a:lstStyle | ./a:txBody/a:lstStyle")
        bp = _xp(el, "./p:txBody/a:bodyPr | ./a:txBody/a:bodyPr")
        self.list_styles += ls
        self.body_prs += bp
        ph = _ph(el)
        self.ph = ph
        if ph is not None:
            lay = _find_ph(slide_ctx.layout_el, ph, by_idx=True)
            mas = _find_ph(slide_ctx.master_el, ph, by_idx=False)
            for inh in (lay, mas):
                if inh is not None:
                    self.list_styles += _xp(inh, "./p:txBody/a:lstStyle")
                    self.body_prs += _xp(inh, "./p:txBody/a:bodyPr")
            st = _txstyle(slide_ctx.master_el, ph["type"])
            if st is not None:
                self.list_styles.append(st)
        if slide_ctx.default_style is not None:
            self.list_styles.append(slide_ctx.default_style)

    def ppr_chain(self, p_el, level):
        chain = []
        own = p_el.find(qn("a:pPr")) if p_el is not None else None
        if own is not None:
            chain.append(own)
        for ls in self.list_styles:
            lv = ls.find(qn(f"a:lvl{level + 1}pPr"))
            if lv is not None:
                chain.append(lv)
        for ls in self.list_styles:
            d = ls.find(qn("a:defPPr"))
            if d is not None:
                chain.append(d)
        return chain

    def body_attr(self, name, default=None):
        for bp in self.body_prs:
            v = bp.get(name)
            if v is not None:
                return v
        return default

    def autofit(self):
        for bp in self.body_prs:
            for c in bp:
                t = _local(c)
                if t == "normAutofit":
                    return "normal", (_int(c.get("fontScale"), 100000) / 100000,
                                      _int(c.get("lnSpcReduction"), 0) / 100000)
                if t == "spAutoFit":
                    return "shape", (1.0, 0.0)
                if t == "noAutofit":
                    return "none", (1.0, 0.0)
        return "none", (1.0, 0.0)


def _first_attr(elems, name):
    for e in elems:
        if e is not None:
            v = e.get(name)
            if v is not None:
                return v
    return None


def _first_child(elems, tag):
    for e in elems:
        if e is not None:
            c = e.find(qn(tag))
            if c is not None:
                return c
    return None


def _spacing(ppr_chain, tag):
    """Line or paragraph spacing: returns ('pct', factor) or ('pts', points)."""
    el = _first_child(ppr_chain, tag)
    if el is None:
        return None
    pct = el.find(qn("a:spcPct"))
    if pct is not None:
        return ("pct", _int(pct.get("val"), 100000) / 100000)
    pts = el.find(qn("a:spcPts"))
    if pts is not None:
        return ("pts", _int(pts.get("val"), 0) / 100)
    return None


# ------------------------------------------------------------------ text


def _text_body(el, chain, theme, is_title):
    """Read paragraphs with effective run properties."""
    tx = _xp(el, "./p:txBody | ./a:txBody")
    if not tx:
        return None
    _, (font_scale, ln_red) = chain.autofit()
    paras = []
    for p_el in tx[0].findall(qn("a:p")):
        ppr = p_el.find(qn("a:pPr"))
        level = _int(ppr.get("lvl"), 0) if ppr is not None else 0
        pchain = chain.ppr_chain(p_el, level)
        defrprs = [x.find(qn("a:defRPr")) for x in pchain]
        runs, chars = [], []
        for child in p_el:
            tag = _local(child)
            if tag not in ("r", "br", "fld"):
                continue
            rpr = child.find(qn("a:rPr"))
            props = [rpr] + defrprs
            sz = _first_attr(props, "sz")
            size = (_int(sz, 1800) / 100.0) if sz else DEFAULT_SIZE_PT
            size = round(size * font_scale, 2)
            latin = _first_child(props, "a:latin")
            font = theme.font(latin.get("typeface") if latin is not None else None)
            if not font:
                font = theme.major if is_title else theme.minor
            bold = _first_attr(props, "b") in ("1", "true")
            color = None
            for pr in props:
                if pr is not None:
                    c = _solid_color(pr, theme)
                    if c:
                        color = c
                        break
            if tag == "br":
                chars.append(("\n", size, font, bold))
                continue
            t = child.find(qn("a:t"))
            text = t.text if t is not None and t.text else ""
            if not text:
                continue
            runs.append({"text": text, "size_pt": size, "font": font, "bold": bold,
                         "color": color})
            chars.extend((ch, size, font, bold) for ch in text)
        text = "".join(c[0] for c in chars)
        end = p_el.find(qn("a:endParaRPr"))
        end_sz = _first_attr([end] + defrprs, "sz")
        def_size = round(((_int(end_sz, 1800) / 100.0) if end_sz else DEFAULT_SIZE_PT)
                         * font_scale, 2)
        algn = _first_attr(pchain, "algn") or "l"
        mar_l = _int(_first_attr(pchain, "marL"), 0) or 0
        paras.append({
            "text": text,
            "level": level,
            "runs": runs,
            "chars": chars,
            "default_size_pt": def_size,
            "align": algn,
            "mar_l": mar_l,
            "line_spacing": _spacing(pchain, "a:lnSpc"),
            "space_before": _spacing(pchain, "a:spcBef"),
            "space_after": _spacing(pchain, "a:spcAft"),
            "ln_reduction": ln_red,
        })
    return paras


def _fit(paras, width, height, chain):
    """Estimate lines and height needed; compare with the box."""
    ins = [
        _int(chain.body_attr("lIns"), DEFAULT_INSETS[0]),
        _int(chain.body_attr("tIns"), DEFAULT_INSETS[1]),
        _int(chain.body_attr("rIns"), DEFAULT_INSETS[2]),
        _int(chain.body_attr("bIns"), DEFAULT_INSETS[3]),
    ]
    wrap = chain.body_attr("wrap", "square") != "none"
    avail_w_pt = (width - ins[0] - ins[2]) / EMU_PER_PT
    total_pt = 0.0
    max_w = 0.0
    widest = 0.0
    n_lines = 0
    for i, p in enumerate(paras):
        line_avail = avail_w_pt - p["mar_l"] / EMU_PER_PT
        m = textfit.measure(p["chars"], line_avail, wrap=wrap)
        p["lines"] = m["lines"]
        n_lines += m["lines"]
        max_w = max(max_w, m["max_line_w_pt"] + p["mar_l"] / EMU_PER_PT)
        widest = max(widest, m["widest_token_pt"])
        ls = p["line_spacing"]
        for s in m["line_sizes"]:
            s = s or p["default_size_pt"]
            if ls and ls[0] == "pts":
                total_pt += ls[1]
            else:
                factor = ls[1] if ls else 1.0
                factor = max(0.5, factor - p["ln_reduction"])
                total_pt += 1.2 * s * factor
        for key in ("space_before", "space_after"):
            sp = p[key]
            if sp and not (key == "space_before" and i == 0):
                total_pt += sp[1] if sp[0] == "pts" else sp[1] * p["default_size_pt"]
    needed = int(total_pt * EMU_PER_PT) + ins[1] + ins[3]
    return {
        "lines": n_lines,
        "needed_height": needed,
        "used_width": int(max_w * EMU_PER_PT) + ins[0] + ins[2],
        "avail_width": int(avail_w_pt * EMU_PER_PT),
        "widest_token": int(widest * EMU_PER_PT),
        "wrap": wrap,
        "insets": ins,
        "overflow_ratio": round(needed / height, 3) if height else None,
    }


def _text_bbox(bbox, fit, anchor, align):
    x, y, w, h = bbox
    uw = min(w, fit["used_width"]) if fit["wrap"] else fit["used_width"]
    uh = fit["needed_height"]
    if align == "ctr":
        tx = x + (w - uw) // 2
    elif align == "r":
        tx = x + w - uw
    else:
        tx = x
    if anchor == "ctr":
        ty = y + (h - uh) // 2
    elif anchor == "b":
        ty = y + h - uh
    else:
        ty = y
    return [int(tx), int(ty), int(uw), int(uh)]


# ------------------------------------------------------------------ images


def image_features(blob: bytes):
    """Cheap statistics that separate chart-like images from photos.

    Returns None if Pillow is not installed or the image cannot be read.
    """
    if len(blob) > MAX_IMAGE_BYTES:
        return None
    try:
        from PIL import Image
    except ImportError:
        return None
    try:
        with Image.open(io.BytesIO(blob)) as im:
            px_size = im.size
            im = im.convert("RGB")
            im.thumbnail((240, 240))
            w, h = im.size
            raw = im.tobytes()
            data = [tuple(raw[i:i + 3]) for i in range(0, len(raw), 3)]
    except Exception:
        return None
    n = len(data) or 1
    quant = {(r >> 4, g >> 4, b >> 4) for r, g, b in data}
    light = sum(1 for r, g, b in data if r > 235 and g > 235 and b > 235) / n

    def dark(p):
        return sum(p) < 3 * 110

    long_rows = 0
    for yy in range(h):
        row = data[yy * w:(yy + 1) * w]
        if sum(1 for p in row if dark(p)) >= 0.6 * w:
            long_rows += 1
    long_cols = 0
    for xx in range(w):
        col = data[xx::w]
        if sum(1 for p in col if dark(p)) >= 0.6 * h:
            long_cols += 1
    return {
        "px": list(px_size),
        "colors": len(quant),
        "light_frac": round(light, 3),
        "axis_rows": long_rows,
        "axis_cols": long_cols,
    }


# ------------------------------------------------------------------ shapes


def _bbox_of(el):
    off = _xp(el, "./p:spPr/a:xfrm/a:off | ./p:xfrm/a:off | ./p:grpSpPr/a:xfrm/a:off")
    ext = _xp(el, "./p:spPr/a:xfrm/a:ext | ./p:xfrm/a:ext | ./p:grpSpPr/a:xfrm/a:ext")
    if not off or not ext:
        return None
    return [_int(off[0].get("x"), 0), _int(off[0].get("y"), 0),
            _int(ext[0].get("cx"), 0), _int(ext[0].get("cy"), 0)]


def _rotation(el):
    x = _xp(el, "./p:spPr/a:xfrm | ./p:xfrm")
    return (_int(x[0].get("rot"), 0) / 60000.0) if x else 0.0


def _rotated_bbox(b, rot):
    if not rot or rot % 180 == 0:
        return b
    if rot % 90 == 0:
        x, y, w, h = b
        cx, cy = x + w / 2, y + h / 2
        return [int(cx - h / 2), int(cy - w / 2), h, w]
    import math

    x, y, w, h = b
    cx, cy = x + w / 2, y + h / 2
    a = math.radians(rot)
    bw = abs(w * math.cos(a)) + abs(h * math.sin(a))
    bh = abs(w * math.sin(a)) + abs(h * math.cos(a))
    return [int(cx - bw / 2), int(cy - bh / 2), int(bw), int(bh)]


def _inherited_bbox(el, slide_ctx):
    """Position for placeholders that inherit their frame from the layout/master."""
    b = _bbox_of(el)
    if b is not None:
        return b
    ph = _ph(el)
    if ph is None:
        return None
    for container, by_idx in ((slide_ctx.layout_el, True), (slide_ctx.master_el, False)):
        inh = _find_ph(container, ph, by_idx=by_idx)
        if inh is not None:
            b = _bbox_of(inh)
            if b is not None:
                return b
    return None


class _SlideCtx:
    def __init__(self, slide, prs_default_style):
        self.layout_el = self.master_el = None
        try:
            self.layout_el = slide.slide_layout._element
            self.master_el = slide.slide_layout.slide_master._element
        except Exception:
            pass
        self.default_style = prs_default_style


def _transform(b, xf):
    """Map a child-space bbox through a group transform list."""
    x, y, w, h = b
    for (ox, oy, sx, sy, cx, cy) in reversed(xf):
        x = ox + (x - cx) * sx
        y = oy + (y - cy) * sy
        w = w * sx
        h = h * sy
    return [int(x), int(y), int(w), int(h)]


def _group_xf(grp_el):
    xfrm = _xp(grp_el, "./p:grpSpPr/a:xfrm")
    if not xfrm:
        return (0, 0, 1.0, 1.0, 0, 0)
    x = xfrm[0]

    def g(tag, a, b):
        e = x.find(qn(tag))
        return (_int(e.get(a), 0), _int(e.get(b), 0)) if e is not None else (0, 0)

    ox, oy = g("a:off", "x", "y")
    cx_, cy_ = g("a:ext", "cx", "cy")
    chx, chy = g("a:chOff", "x", "y")
    chw, chh = g("a:chExt", "cx", "cy")
    sx = cx_ / chw if chw else 1.0
    sy = cy_ / chh if chh else 1.0
    return (ox, oy, sx, sy, chx, chy)


def _iter_elements(sptree, xf=()):
    """Yield (element, group_transforms) for every leaf shape."""
    for el in sptree:
        tag = _local(el)
        if tag == "grpSp":
            yield from _iter_elements(el, xf + (_group_xf(el),))
        elif tag in ("sp", "pic", "graphicFrame", "cxnSp", "contentPart"):
            yield el, xf
        elif tag == "AlternateContent":
            choice = _xp(el, "./mc:Choice/*[1]") or _xp(el, "./mc:Fallback/*[1]")
            if choice:
                yield choice[0], xf


def _shape_record(el, xf, slide, slide_ctx, theme, shape_lookup):
    tag = _local(el)
    cnv = _xp(el, "./*[1]/p:cNvPr")
    sid = _int(cnv[0].get("id")) if cnv else None
    name = cnv[0].get("name", "") if cnv else ""
    descr = cnv[0].get("descr", "") if cnv else ""
    ph = _ph(el)
    raw = _inherited_bbox(el, slide_ctx)
    rot = _rotation(el)
    rec = {
        "id": sid,
        "name": name,
        "kind": "other",
        "placeholder": ph["type"] if ph else None,
        "bbox": None,
        "rotation": rot,
        "in_group": bool(xf),
        "text": "",
    }
    if raw is not None:
        b = _transform(raw, xf) if xf else raw
        rec["bbox"] = _rotated_bbox(b, rot)

    if tag == "cxnSp":
        rec["kind"] = "line"
        return rec
    if tag == "pic":
        rec["kind"] = "picture"
        rec["alt_text"] = descr
        rec["image"] = _image_info(el, slide)
        return rec
    if tag == "contentPart":
        rec["kind"] = "ink"
        return rec
    if tag == "graphicFrame":
        uri = (_xp(el, "./a:graphic/a:graphicData/@uri") or [""])[0]
        if uri == URI_CHART:
            rec["kind"] = "chart"
            rec["chart"] = _chart_info(shape_lookup.get(sid), theme)
        elif "chartex" in uri.lower():
            rec["kind"] = "chart"
            rec["chart"] = {"type": "chartex", "series": [], "min_font_pt": None}
        elif uri == URI_TABLE:
            rec["kind"] = "table"
            rec["table"] = _table_info(el, theme, slide_ctx)
            rec["text"] = "\n".join(c["text"] for c in rec["table"]["cells"])
            if rec["bbox"]:
                # Rows grow to fit their text, so the drawn table can be taller.
                rec["extent"] = rec["bbox"][:3] + [max(rec["bbox"][3],
                                                       rec["table"]["est_height"])]
        elif uri == URI_DIAGRAM:
            rec["kind"] = "diagram"
        else:
            rec["kind"] = "object"
        return rec

    # p:sp: autoshape, text box or placeholder
    prst = (_xp(el, "./p:spPr/a:prstGeom/@prst") or [None])[0]
    rec["geometry"] = prst
    rec["fill"] = _shape_fill(el, theme)
    if prst in ("line", "straightConnector1"):
        rec["kind"] = "line"
        return rec
    chain = StyleChain(el, slide_ctx)
    is_title = bool(ph and ph["type"] in TITLE_TYPES)
    paras = _text_body(el, chain, theme, is_title)
    text = "\n".join(p["text"] for p in paras) if paras else ""
    rec["text"] = text
    if not text.strip():
        rec["kind"] = "placeholder" if ph else "shape"
        rec["empty_placeholder"] = bool(ph) and ph["type"] not in FOOTER_TYPES
        return rec
    rec["kind"] = "text"
    autofit, (scale, _) = chain.autofit()
    rec["autofit"] = autofit
    rec["font_scale"] = scale
    rec["anchor"] = chain.body_attr("anchor", "t")
    if rec["bbox"]:
        fit = _fit(paras, rec["bbox"][2], rec["bbox"][3], chain)
        rec["fit"] = fit
        rec["text_bbox"] = _text_bbox(rec["bbox"], fit, rec["anchor"],
                                      paras[0]["align"] if paras else "l")
        if autofit == "shape" and fit["needed_height"] > rec["bbox"][3]:
            # The box grows to fit its text when rendered.
            rec["extent"] = rec["bbox"][:3] + [fit["needed_height"]]
    rec["paragraphs"] = [{
        "text": p["text"],
        "level": p["level"],
        "lines": p.get("lines"),
        "runs": p["runs"],
        "sizes": sorted({r["size_pt"] for r in p["runs"]}),
    } for p in paras]
    return rec


def _shape_fill(el, theme):
    sppr = _xp(el, "./p:spPr")
    if sppr:
        if sppr[0].find(qn("a:noFill")) is not None:
            return None
        c = _solid_color(sppr[0], theme)
        if c:
            return c
    ref = _xp(el, "./p:style/a:fillRef")
    if ref and ref[0].get("idx") not in (None, "0") and len(ref[0]):
        return theme.color(ref[0][0])
    return None


def _image_info(el, slide):
    info = {"content_type": None, "ext": None, "features": None}
    rid = (_xp(el, "./p:blipFill/a:blip/@r:embed") or [None])[0]
    if not rid:
        return info
    try:
        part = slide.part.related_part(rid)
        info["content_type"] = part.content_type
        info["ext"] = Path(str(part.partname)).suffix.lstrip(".").lower()
        if info["ext"] in ("png", "jpg", "jpeg", "gif", "bmp", "tif", "tiff", "webp"):
            info["features"] = image_features(part.blob)
    except Exception:
        pass
    return info


def _chart_info(shape, theme):
    info = {"type": None, "series": [], "min_font_pt": None, "has_title": False}
    if shape is None:
        return info
    try:
        chart = shape.chart
        info["type"] = str(chart.chart_type).split(".")[-1].split(" ")[0]
        info["has_title"] = bool(chart.has_title)
        cs = chart._chartSpace
        sizes = [int(v) / 100 for v in _xp(cs, ".//a:defRPr/@sz | .//a:rPr/@sz")
                 if str(v).isdigit()]
        info["min_font_pt"] = min(sizes) if sizes else None
        for i, ser in enumerate(_xp(cs, ".//c:ser")):
            name = "".join(_xp(ser, "./c:tx//c:v/text()")) or f"Series {i + 1}"
            color = None
            sppr = ser.find(qn("c:spPr"))
            if sppr is not None:
                color = _solid_color(sppr, theme)
            pt_colors = []
            for dpt in _xp(ser, "./c:dPt/c:spPr"):
                pc = _solid_color(dpt, theme)
                if pc:
                    pt_colors.append(pc)
            info["series"].append({"name": name, "color": color, "point_colors": pt_colors})
        info["vary_colors"] = bool(_xp(cs, ".//c:varyColors[@val='1']"))
    except Exception as exc:  # unusual chart XML must not stop the review
        info["error"] = str(exc)[:120]
    return info


def _table_info(el, theme, slide_ctx):
    cells = []
    fills = []
    tbl = _xp(el, "./a:graphic/a:graphicData/a:tbl")
    if not tbl:
        return {"cells": cells, "fills": fills, "est_height": 0, "has_numbers": False}
    tbl = tbl[0]
    col_w = [_int(g.get("w"), 0) for g in _xp(tbl, "./a:tblGrid/a:gridCol")]
    default_style = slide_ctx.default_style
    est_h = 0
    for r, tr in enumerate(_xp(tbl, "./a:tr")):
        row_h = _int(tr.get("h"), 0)
        need_row = row_h
        for c, tc in enumerate(_xp(tr, "./a:tc")):
            chain = _TableCellChain(tc, default_style)
            paras = _text_body(tc, chain, theme, False) or []
            text = "\n".join(p["text"] for p in paras)
            sizes = [run["size_pt"] for p in paras for run in p["runs"]]
            fonts = [run["font"] for p in paras for run in p["runs"]]
            colors = [run["color"] for p in paras for run in p["runs"] if run["color"]]
            tcpr = tc.find(qn("a:tcPr"))
            fill = _solid_color(tcpr, theme) if tcpr is not None else None
            if fill:
                fills.append(fill)
            w = col_w[c] if c < len(col_w) else 0
            if w and paras:
                fit = _fit(paras, w, row_h or 1, chain)
                need_row = max(need_row, fit["needed_height"])
            cells.append({"r": r, "c": c, "text": text, "sizes": sorted(set(sizes)),
                          "fonts": sorted(set(f for f in fonts if f)),
                          "colors": colors, "runs": [run for p in paras for run in p["runs"]]})
        est_h += need_row
    has_numbers = any(re.search(r"\d", c["text"]) for c in cells if c["r"] > 0)
    return {"cells": cells, "fills": fills, "est_height": est_h, "has_numbers": has_numbers}


class _TableCellChain(StyleChain):
    def __init__(self, tc, default_style):  # table cells have no placeholder chain
        self.list_styles = _xp(tc, "./a:txBody/a:lstStyle")
        if default_style is not None:
            self.list_styles.append(default_style)
        tcpr = tc.find(qn("a:tcPr"))
        self.body_prs = []
        self.ph = None
        self._tcpr = tcpr

    def body_attr(self, name, default=None):
        mapping = {"lIns": ("marL", 91440), "rIns": ("marR", 91440),
                   "tIns": ("marT", 45720), "bIns": ("marB", 45720)}
        if name in mapping and self._tcpr is not None:
            v = self._tcpr.get(mapping[name][0])
            if v is not None:
                return v
        if name in mapping:
            return mapping[name][1]
        return default

    def autofit(self):
        return "none", (1.0, 0.0)


# ------------------------------------------------------------------ roles


def _assign_roles(shapes, slide_h):
    title = None
    for s in shapes:
        if s["kind"] != "text":
            continue
        ph = s["placeholder"]
        if ph in TITLE_TYPES:
            s["role"] = "title"
            if title is None:
                title = s
        elif ph == "subTitle":
            s["role"] = "subtitle"
        elif ph in FOOTER_TYPES:
            s["role"] = "footer"
        elif SOURCE_RE.match(s["text"]) or (SOURCE_ANYWHERE_RE.search(s["text"])
                                             and len(s["text"]) < 300
                                             and s["bbox"] and s["bbox"][1] > slide_h * 0.6):
            s["role"] = "source"
        elif s["bbox"] and s["bbox"][1] >= slide_h * 0.9 and len(s["text"]) <= 60:
            s["role"] = "footer"
        elif (s["bbox"] and s["bbox"][1] + s["bbox"][3] <= slide_h * 0.12
              and len(s["text"]) <= 40 and len(s["text"].split()) <= 6):
            s["role"] = "label"  # sticker, tracker, logo text above the title
        else:
            s["role"] = "body"
    if title is None:
        # Infer a title from a big text box near the top (decks made without placeholders).
        best = None
        for s in shapes:
            if s["kind"] != "text" or s.get("role") != "body" or not s["bbox"]:
                continue
            if s["bbox"][1] > slide_h * 0.22 or len(s["text"]) > 250:
                continue
            size = max((r["size_pt"] for p in s["paragraphs"] for r in p["runs"]), default=0)
            if size < 20:
                continue
            key = (size, -s["bbox"][1])
            if best is None or key > best[0]:
                best = (key, s)
        if best:
            title = best[1]
            title["role"] = "title"
            title["title_inferred"] = True
    for s in shapes:
        s.setdefault("role", s["kind"])
    return title


# ------------------------------------------------------------------ entry point


def _check_archive(path: Path, limit: int):
    try:
        with zipfile.ZipFile(path) as z:
            total = sum(i.file_size for i in z.infolist())
    except zipfile.BadZipFile as exc:
        raise ValueError(f"not a .pptx (zip) file: {exc}") from None
    if total > limit:
        raise ValueError(f"deck unzips to {total / 1e6:.0f} MB, too large to inspect safely "
                         f"(limit {limit / 1e6:.0f} MB)")


def inspect_pptx(path, max_unzipped: int = MAX_UNZIPPED) -> dict:
    """Return a JSON-serialisable description of the deck at `path`."""
    path = Path(path)
    _check_archive(path, max_unzipped)
    prs = Presentation(str(path))
    slide_w = prs.slide_width or 9144000
    slide_h = prs.slide_height or 6858000
    dts = _xp(prs.part._element, "./p:defaultTextStyle")
    default_style = dts[0] if dts else None
    themes = {}
    deck = {
        "file": path.name,
        "slide_width": int(slide_w),
        "slide_height": int(slide_h),
        "slides": [],
        "warnings": [],
    }
    first_theme = None
    for idx, slide in enumerate(prs.slides, start=1):
        try:
            master = slide.slide_layout.slide_master
            key = id(master)
            if key not in themes:
                themes[key] = Theme(master)
            theme = themes[key]
        except Exception:
            theme = Theme.__new__(Theme)
            Theme.__init__(theme, prs.slide_master)
        first_theme = first_theme or theme
        ctx = _SlideCtx(slide, default_style)
        lookup = {}
        try:
            for shp in slide.shapes:
                lookup[shp.shape_id] = shp
                if shp.shape_type == MSO_SHAPE_TYPE.GROUP:
                    stack = list(shp.shapes)
                    while stack:
                        sub = stack.pop()
                        lookup[sub.shape_id] = sub
                        if sub.shape_type == MSO_SHAPE_TYPE.GROUP:
                            stack.extend(sub.shapes)
        except Exception:
            pass
        shapes = []
        sptree = slide.shapes._spTree
        for order, (el, xf) in enumerate(_iter_elements(sptree)):
            try:
                rec = _shape_record(el, xf, slide, ctx, theme, lookup)
            except Exception as exc:  # keep going on odd shapes
                deck["warnings"].append(f"slide {idx}: could not read a shape ({exc})")
                continue
            rec["z"] = order
            if not rec["name"]:
                rec["name"] = f"{rec['kind']} {rec['id']}"
            rec.setdefault("extent", rec["bbox"])
            shapes.append(rec)
        title = _assign_roles(shapes, slide_h)
        notes = ""
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text or ""
        except Exception:
            pass
        try:
            layout_name = slide.slide_layout.name
        except Exception:
            layout_name = None
        words = sum(len(s["text"].split()) for s in shapes
                    if s.get("role") in ("body", "title", "subtitle", "source") or
                    s["kind"] == "table")
        deck["slides"].append({
            "index": idx,
            "layout": layout_name,
            "hidden": slide._element.get("show") in ("0", "false"),
            "title": title["text"].strip() if title else None,
            "title_id": title["id"] if title else None,
            "words": words,
            "notes": notes,
            "shapes": shapes,
        })
    if first_theme is not None:
        deck["theme"] = {"major_font": first_theme.major, "minor_font": first_theme.minor,
                         "colors": first_theme.colors}
    return deck
