"""Safe automatic fixes. Writes a new file; the input deck is never modified.

Only layout and size change. Words, numbers, chart data and slide order never
change. The one addition is a visible "Source: [SOURCE NEEDED]" line on exhibit
slides that have none, so a person fills in the real source.

Fixes:
  font_below_floor  raise runs under the floor to the floor (body 12, source 10,
                    table 10 pt)
  off_slide         move a shape that sticks out back inside the slide
                    (shapes entirely off the slide are left alone: they may be parked)
  missing_source    add the source placeholder line under the exhibit
  misaligned        snap a near-miss edge to its neighbour's edge

Shapes that are part of an exhibit (see checks.exhibit_parts: charts, tables,
pictures, drawn charts and grids, groups, connectors, shapes named sw:*, and
anything sitting on an exhibit) are never moved, resized or restyled; their
findings are logged as "left as is". The one exception is raising the text of
a plain native table to the floor, which changes no shape's position or size.
"""

from __future__ import annotations

import math
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from . import checks as C
from .inspect import inspect_pptx

EMU_IN = 914400
SOURCE_PLACEHOLDER = {"en": "Source: [SOURCE NEEDED]", "ko": "출처: [SOURCE NEEDED]",
                      "ja": "出典: [SOURCE NEEDED]"}
FIXABLE = ("font_below_floor", "off_slide", "missing_source", "misaligned")


def _shape_index(slide):
    """{shape_id: shape} for top-level shapes (group children are not moved)."""
    return {shp.shape_id: shp for shp in slide.shapes}


def _record(deck, slide_idx, shape_id):
    for s in deck["slides"][slide_idx - 1]["shapes"]:
        if s["id"] == shape_id:
            return s
    return None


def _language(deck):
    text = " ".join(s.get("title") or "" for s in deck["slides"])
    if re.search(r"[가-힣]", text):
        return "ko"
    if re.search(r"[぀-ヿ]", text):
        return "ja"
    return "en"


def _text_runs(txbody_el):
    """(paragraph text, [run elements with text]) per paragraph, in document order."""
    for p in txbody_el.findall(qn("a:p")):
        runs = [el for el in p if el.tag in (qn("a:r"), qn("a:fld"))
                and (el.findtext(qn("a:t")) or "")]
        text = "".join(el.findtext(qn("a:t")) or "" for el in p
                       if el.tag in (qn("a:r"), qn("a:fld")))
        yield text, runs


def _raise_runs(txbody_el, sizes, floor_for_para, scale=1.0):
    """Set sz on runs whose effective size is under the floor. Returns count."""
    changed = 0
    flat = iter(sizes)
    for ptext, runs in _text_runs(txbody_el):
        floor = floor_for_para(ptext)
        for r in runs:
            size = next(flat, None)
            if size is None or not (r.findtext(qn("a:t")) or "").strip():
                continue
            if size < floor - 0.05:
                rpr = r.find(qn("a:rPr"))
                if rpr is None:
                    rpr = r.makeelement(qn("a:rPr"), {})
                    r.insert(0, rpr)
                rpr.set("sz", str(int(math.ceil(floor / scale * 100))))
                changed += 1
    return changed


def _fix_font(slide, rec, issue, cfg):
    shp = _shape_index(slide).get(rec["id"])
    if shp is None:
        return None
    if rec["kind"] == "text":
        sizes = [r["size_pt"] for p in rec["paragraphs"] for r in p["runs"]]
        role = rec.get("role")

        def floor_for(ptext):
            if role in ("source", "footer", "label") or C.SOURCE_LINE.match(ptext):
                return cfg["source_min_pt"]
            return cfg["body_min_pt"]

        txbody = shp._element.find(qn("p:txBody"))
        n = _raise_runs(txbody, sizes, floor_for, rec.get("font_scale") or 1.0)
    elif rec["kind"] == "table":
        n = 0
        cells = iter(rec["table"]["cells"])
        for tr in shp._element.iter(qn("a:tr")):
            for tc in tr.findall(qn("a:tc")):
                cell = next(cells, None)
                txbody = tc.find(qn("a:txBody"))
                if cell is None or txbody is None:
                    continue
                n += _raise_runs(txbody, [r["size_pt"] for r in cell["runs"]],
                                 lambda _t: cfg["table_min_pt"])
    else:
        return None
    if not n:
        return None
    return f"raised {n} text run(s) to the {issue.get('floor_pt', 0):g} pt floor"


def _set_frame(shp, rec, left=None, top=None, width=None):
    """Move/resize; placeholders that inherit their frame get an explicit one."""
    x, y, w, h = rec["bbox"]
    if shp.left is None or getattr(shp, "is_placeholder", False):
        shp.left, shp.top, shp.width, shp.height = Emu(x), Emu(y), Emu(w), Emu(h)
    if left is not None:
        shp.left = Emu(int(left))
    if top is not None:
        shp.top = Emu(int(top))
    if width is not None:
        shp.width = Emu(int(width))


def _fix_off_slide(deck, slide, rec, issue):
    if rec.get("in_group") or issue.get("fully_off") or rec.get("rotation"):
        return None
    shp = _shape_index(slide).get(rec["id"])
    if shp is None:
        return None
    W, H = deck["slide_width"], deck["slide_height"]
    x, y, w, _ = rec["bbox"]
    h = (rec.get("extent") or rec["bbox"])[3]
    if w > W or h > H:
        return None
    nx = min(max(x, 0), W - w)
    ny = min(max(y, 0), H - h)
    if (nx, ny) == (x, y):
        return None
    _set_frame(shp, rec, left=nx, top=ny)
    return f"moved inside the slide by ({(nx - x) / EMU_IN:+.2f} in, {(ny - y) / EMU_IN:+.2f} in)"


def _fix_alignment(deck, slide, rec, issue):
    if rec.get("in_group"):
        return None
    anchor = _record(deck, issue["slide"], issue.get("anchor_id"))
    shp = _shape_index(slide).get(rec["id"])
    if shp is None or anchor is None or not anchor.get("bbox"):
        return None
    x, y, w, _ = rec["bbox"]
    ax, ay, aw, _ = anchor["bbox"]
    edge = issue.get("edge")
    if edge == "left":
        if rec["kind"] == "picture":
            _set_frame(shp, rec, left=ax)  # move, keep the aspect ratio
        else:
            _set_frame(shp, rec, left=ax, width=x + w - ax)  # keep the right edge
    elif edge == "top":
        _set_frame(shp, rec, top=ay)
    elif edge == "right":
        _set_frame(shp, rec, width=ax + aw - x)
    else:
        return None
    return f"snapped its {edge} edge to \"{anchor['name']}\""


def _fix_source(deck, slide, slide_rec, lang, exhibit_ids=frozenset()):
    W, H = deck["slide_width"], deck["slide_height"]
    margin = int(0.5 * EMU_IN)
    title = next((s for s in slide_rec["shapes"] if s["id"] == slide_rec.get("title_id")),
                 None)
    left = title["bbox"][0] if title and title.get("bbox") else margin
    left = max(int(0.2 * EMU_IN), min(left, W // 4))
    right = title["bbox"][0] + title["bbox"][2] if title and title.get("bbox") else W - left
    right = min(W - int(0.2 * EMU_IN), max(right, W * 3 // 4))
    height = int(0.32 * EMU_IN)
    bottoms = [s["extent"][1] + s["extent"][3] for s in slide_rec["shapes"]
               if s.get("extent") and (s["kind"] in ("chart", "table", "picture")
                                       or s["id"] in exhibit_ids)]
    top = max(bottoms) + int(0.05 * EMU_IN) if bottoms else H - margin
    if top + height > H - int(0.1 * EMU_IN):
        top = H - height - int(0.1 * EMU_IN)
    box = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(right - left), Emu(height))
    box.name = "Source (added by deck-review)"
    tf = box.text_frame
    tf.word_wrap = True
    run = tf.paragraphs[0].add_run()
    run.text = SOURCE_PLACEHOLDER[lang]
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    return "added a 'Source: [SOURCE NEEDED]' line under the exhibit"


def fix_pptx(src, dst, issues: list | None = None, config: dict | None = None) -> list[dict]:
    """Apply safe fixes for `issues` (found automatically if None) and save to dst.

    Returns one record per fix: {check, slide, shape, shape_id, action, ok}.
    """
    src, dst = Path(src), Path(dst)
    if src.resolve() == dst.resolve():
        raise ValueError("fix_pptx writes a new file; dst must differ from src")
    cfg = dict(C.DEFAULTS)
    if config:
        cfg.update(config)
    deck = inspect_pptx(src)
    if issues is None:
        issues = C.run_checks(deck, cfg)
    prs = Presentation(str(src))
    slides = list(prs.slides)
    lang = _language(deck)
    applied = []
    sourced = set()
    parts = {}  # slide -> (data ids, frame ids), see checks.exhibit_parts
    for issue in issues:
        check, idx = issue["check"], issue.get("slide")
        if check not in FIXABLE or not issue.get("auto_fix") or idx is None:
            continue
        slide = slides[idx - 1]
        rec = _record(deck, idx, issue.get("shape_id")) if issue.get("shape_id") else None
        if idx not in parts:
            parts[idx] = C.exhibit_parts(deck, deck["slides"][idx - 1])
        data, frames = parts[idx]
        if rec is not None and check != "missing_source" and (
                rec["id"] in data or (rec["id"] in frames and not (
                    check == "font_below_floor" and rec["kind"] == "table"))):
            # Part of an exhibit: its position and size carry the data, so it is
            # reported, never moved, resized or restyled. (A plain native table
            # may have its text raised to the floor; that changes no geometry.)
            applied.append({"check": check, "slide": idx, "shape": issue.get("shape"),
                            "shape_id": issue.get("shape_id"),
                            "action": "left as is: part of an exhibit; check by hand",
                            "ok": False})
            continue
        action = None
        try:
            if check == "missing_source" and idx not in sourced:
                action = _fix_source(deck, slide, deck["slides"][idx - 1], lang,
                                     data | frames)
                sourced.add(idx)
            elif rec is None:
                continue
            elif check == "font_below_floor":
                action = _fix_font(slide, rec, issue, cfg)
            elif check == "off_slide":
                action = _fix_off_slide(deck, slide, rec, issue)
            elif check == "misaligned":
                action = _fix_alignment(deck, slide, rec, issue)
        except Exception as exc:  # one odd shape must not stop the other fixes
            action = None
            applied.append({"check": check, "slide": idx, "shape": issue.get("shape"),
                            "shape_id": issue.get("shape_id"),
                            "action": f"skipped ({exc})", "ok": False})
        if action:
            applied.append({"check": check, "slide": idx, "shape": issue.get("shape"),
                            "shape_id": issue.get("shape_id"), "action": action, "ok": True})
    dst.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(dst))
    return applied
