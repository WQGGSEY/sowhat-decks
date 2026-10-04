"""Build before.pptx: a deliberately weak draft of the Q3 board update.

Same sample data as the finished deck (../inputs/), with four planted problems
that deck-review should catch:

1. topic-label titles ("Bookings performance", "Next steps"), copied from
   ../../storyline-tests/control-topic-titles/storyline.md
2. charts pasted as pictures (PNG) instead of native, editable charts
3. tiny fonts (8-11 pt body and table text)
4. no source lines, and a "TBD" left in the forecast

Run from the repository root:

    python3 examples/03-review-before-after/before/make_before.py

Needs python-pptx and Pillow. Writes before.pptx next to this script.
"""
from __future__ import annotations

import io
import pathlib

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt

HERE = pathlib.Path(__file__).resolve().parent
OUT = HERE / "before.pptx"

FONT_PATHS = ["/System/Library/Fonts/Supplemental/Arial.ttf",
              "/Library/Fonts/Arial.ttf",
              "C:/Windows/Fonts/arial.ttf",
              "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"]
COLORS = [(68, 114, 196), (237, 125, 49), (165, 165, 165)]  # spreadsheet-default blue, orange, grey


def font(size: int):
    for p in FONT_PATHS:
        if pathlib.Path(p).is_file():
            return ImageFont.truetype(p, size)
    return ImageFont.load_default()


def chart_png(title: str, categories: list[str], series: dict[str, list[float]],
              kind: str = "column", y_max: float | None = None) -> io.BytesIO:
    """A spreadsheet-style chart drawn as a picture: legend, gridlines, small labels."""
    w, h = 1200, 675
    img = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(img)
    f_title, f_lab = font(30), font(20)
    d.text((w // 2, 30), title, fill=(64, 64, 64), font=f_title, anchor="mt")
    left, right, top, bottom = 110, w - 60, 100, h - 140
    vals = [v for s in series.values() for v in s]
    y_max = y_max or max(vals) * 1.15
    for i in range(6):  # gridlines and axis labels
        y = bottom - (bottom - top) * i / 5
        d.line([(left, y), (right, y)], fill=(217, 217, 217), width=2)
        d.text((left - 12, y), f"{y_max * i / 5:,.0f}", fill=(89, 89, 89), font=f_lab, anchor="rm")
    d.line([(left, bottom), (right, bottom)], fill=(40, 40, 40), width=6)  # black axes, older spreadsheet default
    d.line([(left, top), (left, bottom)], fill=(40, 40, 40), width=6)
    n, k = len(categories), len(series)
    slot = (right - left) / n
    for ci, cat in enumerate(categories):
        cx = left + slot * (ci + 0.5)
        d.text((cx, bottom + 14), cat, fill=(89, 89, 89), font=f_lab, anchor="mt")
    for si, (name, values) in enumerate(series.items()):
        color = COLORS[si % len(COLORS)]
        pts = []
        for ci, v in enumerate(values):
            cx = left + slot * (ci + 0.5)
            y = bottom - (bottom - top) * v / y_max
            if kind == "column":
                bw = slot * 0.7 / k
                x0 = cx - slot * 0.35 + bw * si
                d.rectangle([x0, y, x0 + bw - 6, bottom], fill=color)
            else:
                pts.append((cx, y))
        if pts:
            d.line(pts, fill=color, width=6)
            for x, y in pts:
                d.ellipse([x - 8, y - 8, x + 8, y + 8], fill=color)
    lx = w // 2 - 160 * k // 2  # legend under the plot
    for si, name in enumerate(series):
        x = lx + si * 220
        d.rectangle([x, h - 60, x + 22, h - 38], fill=COLORS[si % len(COLORS)])
        d.text((x + 32, h - 49), name, fill=(89, 89, 89), font=f_lab, anchor="lm")
    buf = io.BytesIO()
    img.quantize(colors=48).save(buf, format="PNG", optimize=True)
    buf.seek(0)
    return buf


def bullets(slide, items: list[str], size: int = 11) -> None:
    body = slide.placeholders[1].text_frame
    body.text = items[0]
    for item in items[1:]:
        body.add_paragraph().text = item
    for p in body.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)


def table(slide, rows: list[list[str]], size: int = 9) -> None:
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(0.8), Inches(1.7),
                                   Inches(11.7), Inches(0.3 * len(rows)))
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = shape.table.cell(r, c)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(size)


def main() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    TITLE, CONTENT, TITLE_ONLY = prs.slide_layouts[0], prs.slide_layouts[1], prs.slide_layouts[5]

    def content(title: str, items: list[str], size: int = 11):
        s = prs.slides.add_slide(CONTENT)
        s.shapes.title.text = title
        bullets(s, items, size)
        return s

    def picture(title: str, png: io.BytesIO):
        s = prs.slides.add_slide(TITLE_ONLY)
        s.shapes.title.text = title
        s.shapes.add_picture(png, Inches(1.5), Inches(1.6), width=Inches(10.3))
        return s

    s = prs.slides.add_slide(TITLE)
    s.shapes.title.text = "Q3 2026 board update"
    s.placeholders[1].text = "Sample Co. (fictional company). All figures are sample data."

    content("Executive summary", [
        "Q3 new ARR $1,150k vs plan $950k",
        "Ending ARR $11,050k vs plan $10,900k",
        "SMB logo churn 3.1% monthly in Q3 (Q1: 1.8%)",
        "Onboarding team still 2 specialists",
        "Proposal to change the Q4 hiring plan",
        "Cash $14.2M"])
    picture("Bookings performance", chart_png(
        "New ARR vs plan ($k), sample data", ["Q1", "Q2", "Q3"],
        {"Actual": [900, 1000, 1150], "Plan": [850, 950, 950]}, y_max=1400))
    s = prs.slides.add_slide(TITLE_ONLY)
    s.shapes.title.text = "Q3 ARR bridge"
    table(s, [["$k", "Q3"], ["Beginning ARR", "10,160"], ["New", "1,150"], ["Expansion", "300"],
              ["Contraction", "-80"], ["Churned", "-480"], ["Ending ARR", "11,050"],
              ["Plan ending ARR", "10,900"]])
    picture("Churn by segment", chart_png(
        "Churned ARR by segment ($k), sample data", ["Q1", "Q2", "Q3"],
        {"SMB": [170, 220, 390], "Mid-market": [70, 80, 90]}, kind="line", y_max=500))
    content("Q4 outlook", [
        "Q4 plan new ARR $1,050k",
        "Year-end ARR plan $12,000k",
        "Year-end ARR forecast: TBD",
        "Churn trend is a risk to the plan"])
    content("Setup status of churned accounts", [
        "61 SMB accounts churned in Q3",
        "44 had not finished setup within 30 days of signing",
        "17 had finished setup"])
    content("Churn by setup status", [
        "Jan-Jun 2026 SMB cohort, 90-day churn",
        "Setup finished within 30 days: 4%",
        "Setup not finished: 19%"])
    picture("Onboarding capacity", chart_png(
        "Setup completed in 30 days (%) and new accounts per specialist, sample data",
        ["Q1", "Q2", "Q3"],
        {"Setup in 30 days %": [68, 61, 55], "Accounts per specialist": [42.5, 50, 60]},
        kind="line", y_max=80))
    content("Proposal: hiring plan changes", [
        "Convert 2 of the 4 Q4 AE requisitions into onboarding specialists",
        "AE fully loaded about $150k; onboarding specialist about $95k",
        "170 SMB accounts signed Apr-Sep have not finished setup",
        "Two support team members can transfer now"], size=10)
    content("Risks", [
        "2027 sales capacity",
        "Churn could keep rising",
        "Hiring timelines"])
    content("Next steps", [
        "Discuss the hiring plan",
        "Update the forecast",
        "Review churn again in January"])
    content("Cash runway", ["Cash $14.2M at end of Q3", "Net burn about $650k a month"])
    s = prs.slides.add_slide(TITLE_ONLY)
    s.shapes.title.text = "Q1-Q3 2026 metrics by quarter"
    table(s, [["$k", "Q1", "Q2", "Q3"],
              ["New ARR", "900", "1,000", "1,150"], ["Plan new ARR", "850", "950", "950"],
              ["Ending ARR", "9,250", "10,160", "11,050"], ["Plan ending ARR", "9,200", "10,050", "10,900"],
              ["Churned ARR, SMB", "170", "220", "390"], ["Churned ARR, mid-market", "70", "80", "90"],
              ["SMB logo churn, % a month", "1.8", "2.2", "3.1"],
              ["Mid-market logo churn, % a month", "0.6", "0.7", "0.6"]], size=8)
    prs.save(OUT)
    print(f"wrote {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
