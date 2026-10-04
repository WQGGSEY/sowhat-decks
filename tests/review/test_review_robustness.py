"""Odd but valid decks must go through the whole pipeline without crashing."""

import io

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

import review_builders as B
from deckreview.checks import run_checks, slide_scores
from deckreview.fix import fix_pptx
from deckreview.inspect import inspect_pptx
from deckreview.report import build_markdown


def _kitchen_sink(prs):
    B.add_cover(prs, "Kitchen sink")
    s = B.add_titled_slide(prs, "Three regions drove growth while two stayed flat")
    data = CategoryChartData()
    data.categories = ["Q1", "Q2", "Q3"]
    for name in ("North", "South", "East"):
        data.add_series(name, (1, 2, 3))
    s.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(0.6), Inches(1.7), Inches(7), Inches(4.5),
                       data)  # default multi-colour palette
    grp = s.shapes.add_group_shape()
    pic = grp.shapes.add_picture(io.BytesIO(B.chart_png_bytes()), Inches(8), Inches(1.7),
                                 Inches(4), Inches(2.5))
    pic.name = "Grouped chart picture"
    grp.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8), Inches(4.5), Inches(1), Inches(1))
    rot = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(12.5), Inches(3), Inches(2),
                             Inches(0.5))
    rot.rotation = 90
    s.notes_slide.notes_text_frame.text = "Say the North number out loud."
    B.add_source(s)

    s = B.add_titled_slide(prs, "Merged cells still read")
    tbl = B.add_table(s, [["A", "B", "C"], ["1", "2", "3"], ["4", "5", "6"]])
    tbl.table.cell(0, 0).merge(tbl.table.cell(0, 1))
    B.add_source(s)

    s = prs.slides.add_slide(prs.slide_layouts[8])  # picture with caption, all empty
    s.shapes.title.text = "Picture placeholder left empty"

    s = B.add_titled_slide(prs, "This slide is hidden from the show")
    s._element.set("show", "0")


def test_kitchen_sink_deck_runs_end_to_end(make_deck, tmp_path):
    path = make_deck(_kitchen_sink)
    deck = inspect_pptx(path)
    issues = run_checks(deck)
    for n, i in enumerate(issues, 1):
        i["n"] = n
    md = build_markdown(deck, issues, slide_scores(deck, issues))
    fixes = fix_pptx(path, tmp_path / "fixed.pptx", issues=issues)
    assert "## Top 5 fixes" in md
    assert deck["slides"][4]["hidden"] is True
    assert deck["slides"][1]["notes"].startswith("Say the North")
    grouped = [s for s in deck["slides"][1]["shapes"] if s["name"] == "Grouped chart picture"]
    assert grouped and grouped[0]["bbox"][0] == Inches(8)
    checks = {i["check"] for i in issues}
    assert "accent_colors" in checks  # three default series colours
    assert "chart_as_image" in checks  # the picture inside the group
    assert all(f["ok"] for f in fixes)


def test_empty_deck_gives_an_empty_review(make_deck):
    deck = inspect_pptx(make_deck(lambda prs: None))
    assert deck["slides"] == []
    assert run_checks(deck) == []
    assert "No automatic findings" in build_markdown(deck, [], {})
