"""Behaviour of individual checks on small purpose-built decks."""

import review_builders as B
from pptx.enum.text import MSO_AUTO_SIZE
from deckreview.checks import run_checks
from deckreview.inspect import inspect_pptx


def _issues(path, check=None):
    found = run_checks(inspect_pptx(path))
    return [i for i in found if check is None or i["check"] == check]


def test_one_misaligned_box_is_reported_once_against_its_neighbours(make_deck):
    def build(prs):
        s = B.add_titled_slide(prs, "Two hires cut early churn by a third")
        B.add_text(s, "First point.", 0.6, 1.7, 12.1, 0.8)
        B.add_text(s, "Second point.", 0.6, 2.7, 12.1, 0.8)
        B.add_text(s, "Third point, nudged.", 0.7, 3.7, 12.0, 0.8, name="Nudged")

    found = _issues(make_deck(build), "misaligned")
    assert [(i["slide"], i["shape"]) for i in found] == [(1, "Nudged")]


def test_three_line_title_is_too_long(make_deck):
    def build(prs):
        B.add_cover(prs, "Market entry")
        s = B.add_titled_slide(prs, (
            "Our analysis of the three candidate markets across seven criteria shows that, "
            "on balance and with some caveats regarding data quality, Vietnam is the most "
            "attractive market for the first launch next year"))
        B.add_text(s, "Body.", 0.6, 1.7, 12.1, 0.8)

    assert _issues(make_deck(build), "title_too_long")


def test_source_line_below_10pt_is_flagged_but_10pt_passes(make_deck):
    def build(prs):
        s = B.add_titled_slide(prs, "Revenue grew every quarter and reached $4.6M in Q4")
        B.add_bar_chart(s)
        B.add_text(s, "Source: Company filings", 0.6, 6.85, 12.1, 0.35, size=8, name="Src8")
        s = B.add_titled_slide(prs, "Revenue grew every quarter and reached $4.6M in Q4 again")
        B.add_bar_chart(s)
        B.add_text(s, "Source: Company filings", 0.6, 6.85, 12.1, 0.35, size=10)

    found = _issues(make_deck(build), "font_below_floor")
    assert [(i["slide"], i["shape"], i["floor_pt"]) for i in found] == [(1, "Src8", 10.0)]


def test_small_table_text_is_flagged(make_deck):
    def build(prs):
        s = B.add_titled_slide(prs, "Vietnam leads the shortlist on four of five criteria")
        B.add_table(s, [["Market", "Users"], ["Vietnam", "38"]], size=8)
        B.add_source(s)

    assert _issues(make_deck(build), "font_below_floor")


def test_data_needed_markers_are_listed(make_deck):
    def build(prs):
        s = B.add_titled_slide(prs, "Payback fell to [DATA NEEDED: CAC payback] months")
        B.add_text(s, "CAC is TBD until the test ends; margin XX% in Q4.", 0.6, 1.7, 12.1, 1)

    found = _issues(make_deck(build), "data_needed")
    assert {i["severity"] for i in found} == {"high"}
    assert len(found) == 2


def test_repeated_title_is_a_duplicate(make_deck):
    def build(prs):
        B.add_cover(prs, "Churn review")
        for _ in range(2):
            s = B.add_titled_slide(prs, "Enterprise churn doubled after the price change")
            B.add_text(s, "Body.", 0.6, 1.7, 12.1, 0.8)

    found = _issues(make_deck(build), "duplicate_title")
    assert [(i["slide"], i["other_slide"]) for i in found] == [(3, 2)]


def test_fixed_size_box_that_spills_is_overflow(make_deck):
    def build(prs):
        s = B.add_titled_slide(prs, "Vietnam scores highest on five of seven criteria")
        box = B.add_text(s, B.LONG_TEXT, 0.6, 1.7, 6.0, 1.5, size=16, name="Fixed")
        box.text_frame.auto_size = MSO_AUTO_SIZE.NONE

    found = _issues(make_deck(build), "text_overflow")
    assert [(i["shape"], i["severity"]) for i in found] == [("Fixed", "high")]


def test_section_divider_may_keep_a_label_title(make_deck):
    def build(prs):
        B.add_cover(prs, "Board update")
        B.add_titled_slide(prs, "Financials")
        s = B.add_titled_slide(prs, "Revenue grew 18% in Q3")
        B.add_text(s, "Body.", 0.6, 1.7, 12.1, 0.8)

    assert not _issues(make_deck(build), "title_not_claim")


def test_shape_inside_a_group_is_placed_in_slide_coordinates(make_deck):
    from pptx.util import Inches

    def build(prs):
        s = B.add_titled_slide(prs, "Two hires cut early churn by a third")
        grp = s.shapes.add_group_shape()
        box = grp.shapes.add_textbox(Inches(11), Inches(3), Inches(4), Inches(1))
        box.name = "Grouped"
        box.text_frame.text = "Runs off the right edge"

    found = _issues(make_deck(build), "off_slide")
    assert [i["shape"] for i in found] == ["Grouped"]


def test_korean_deck_titles_are_judged(make_deck):
    def build(prs):
        B.add_cover(prs, "3분기 이사회 보고")
        s = B.add_titled_slide(prs, "시장 현황")
        B.add_text(s, "온라인 학습 시장은 연 14% 성장했다.", 0.6, 1.7, 12.1, 0.8)
        s = B.add_titled_slide(prs, "이탈한 SMB 고객 대부분은 가입 30일 안에 설정을 끝내지 못했다")
        B.add_text(s, "설정 완료율 42%", 0.6, 1.7, 12.1, 0.8)

    found = _issues(make_deck(build), "title_not_claim")
    assert [i["slide"] for i in found] == [2]


def test_short_furniture_text_at_the_edges_uses_the_10pt_floor(make_deck):
    def build(prs):
        s = B.add_titled_slide(prs, "Revenue grew 39% to $1.04B in 2025")
        B.add_text(s, "DRAFT", 12.0, 0.15, 0.8, 0.3, size=10, name="Sticker")
        B.add_text(s, "2025 results", 0.6, 0.15, 6.0, 0.3, size=11, name="Tracker")
        B.add_text(s, "Confidential", 0.6, 7.05, 3.0, 0.3, size=8, name="Tiny footer")
        B.add_text(s, "Body text.", 0.6, 1.7, 12.1, 0.8, size=16)

    found = _issues(make_deck(build), "font_below_floor")
    assert [i["shape"] for i in found] == ["Tiny footer"]
    assert not _issues(make_deck(build), "misaligned")


def test_label_title_in_the_appendix_is_only_a_low_issue(make_deck):
    def build(prs):
        B.add_cover(prs, "Investor update")
        s = B.add_titled_slide(prs, "Revenue grew 39% to $1.04B in 2025")
        B.add_text(s, "Body text.", 0.6, 1.7, 12.1, 0.8, size=16)
        B.add_titled_slide(prs, "Appendix")
        s = B.add_titled_slide(prs, "Key figures FY2023-FY2025")
        B.add_table(s, [["Year", "Revenue"], ["2025", "1,040"]])
        B.add_source(s)

    found = _issues(make_deck(build), "title_not_claim")
    assert [(i["slide"], i["severity"]) for i in found] == [(4, "low")]


def test_empty_placeholders_are_reported_once_per_slide(make_deck):
    def build(prs):
        prs.slides.add_slide(prs.slide_layouts[3])  # title + two empty content boxes

    found = _issues(make_deck(build), "empty_placeholder")
    assert len(found) == 1 and "3 empty placeholders" in found[0]["message"]


def test_unnamed_shapes_get_a_readable_name(make_deck):
    def build(prs):
        s = B.add_titled_slide(prs, "Customers cancel over missing integrations")
        a = B.add_text(s, "One.", 0.6, 1.7, 6.0, 1.0)
        b = B.add_text(s, "Two.", 0.7, 1.8, 6.0, 1.0)
        a.name = b.name = ""

    found = _issues(make_deck(build), "overlap")
    assert found and '""' not in found[0]["message"]


def test_sample_data_label_counts_as_provenance_for_an_exhibit(make_deck):
    def build(prs):
        s = B.add_titled_slide(prs, "Revenue grew every quarter and reached $4.6M in Q4")
        B.add_bar_chart(s)
        B.add_text(s, "Sample data", 0.6, 7.0, 10.0, 0.2, size=10)
        s = B.add_titled_slide(prs, "Revenue grew every quarter and reached $4.6M in Q4 too")
        B.add_bar_chart(s)
        B.add_text(s, "Confidential", 0.6, 7.0, 10.0, 0.2, size=10)

    found = _issues(make_deck(build), "missing_source")
    assert [i["slide"] for i in found] == [2]


def test_title_is_inferred_on_decks_without_placeholders(make_deck):
    def build(prs):
        for title in ("Quarterly review", "Pipeline overview"):
            s = prs.slides.add_slide(prs.slide_layouts[B.LAYOUT_BLANK])
            B.add_text(s, title, 0.6, 0.4, 12.1, 0.9, size=28)
            B.add_text(s, "Body text for the slide.", 0.6, 1.7, 12.1, 0.8, size=16)

    path = make_deck(build)
    assert [s["title"] for s in inspect_pptx(path)["slides"]] == ["Quarterly review",
                                                                  "Pipeline overview"]
    assert [i["slide"] for i in _issues(path, "title_not_claim")] == [2]


def test_section_header_with_a_subtitle_is_a_divider(make_deck):
    def build(prs):
        B.add_cover(prs, "Board update")
        s = prs.slides.add_slide(prs.slide_layouts[2])  # "Section Header"
        s.shapes.title.text = "Where did growth come from?"
        s.placeholders[1].text = "By region and product"
        s = B.add_titled_slide(prs, "Region B grew three times faster than the rest")
        B.add_text(s, "Body.", 0.6, 1.7, 12.1, 0.8)

    assert not _issues(make_deck(build), "title_not_claim")
