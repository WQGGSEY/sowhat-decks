"""The safe fix must never move or resize a shape that is part of an exhibit.

In a drawn chart (Mekko, Gantt, harvey balls) the position and size of each
shape are the data, so snapping near-miss edges or nudging shapes changes what
the chart says. Issue #21.
"""

import os
from pathlib import Path

import pytest
import review_builders as B
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from deckreview.checks import run_checks
from deckreview.fix import fix_pptx
from deckreview.inspect import inspect_pptx

REPO = Path(__file__).resolve().parents[2]


def _geometry(path):
    """{(slide, shape path): (name, left, top, width, height)} incl. group members."""
    out = {}

    def walk(shapes, slide_no, prefix):
        for shp in shapes:
            key = (slide_no, f"{prefix}{shp.shape_id}")
            out[key] = (shp.name, shp.left, shp.top, shp.width, shp.height)
            if shp.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(shp.shapes, slide_no, f"{prefix}{shp.shape_id}/")

    for n, slide in enumerate(Presentation(str(path)).slides, 1):
        walk(slide.shapes, n, "")
    return out


def _changed(src, dst):
    before, after = _geometry(src), _geometry(dst)
    return {before[k][0] for k in before if k in after and before[k][1:] != after[k][1:]}


def _pro_showcase():
    candidates = [os.environ.get("SOWHAT_DECKS_PRO", "")]
    for parent in (REPO.parent, REPO.parent.parent):
        candidates.append(str(parent / "sowhat-decks-pro"))
    for c in candidates:
        deck = Path(c) / "examples" / "11-pro-exhibits-showcase" / "deck.pptx"
        if c and deck.exists():
            return deck
    return None


def test_drawn_exhibits_keep_their_geometry_and_safe_shapes_are_still_fixed(decks, tmp_path):
    dst = tmp_path / "drawn.fixed.pptx"
    fix_pptx(decks["drawn"], dst)
    changed = _changed(decks["drawn"], dst)
    assert changed - B.DRAWN_SAFE_SHAPES == set()
    assert changed == {"Off-slide note", "Nudged note"}  # moved in, snapped
    tiny = [s for s in Presentation(str(dst)).slides[3].shapes if s.name == "Tiny note"][0]
    assert tiny.text_frame.paragraphs[0].runs[0].font.size.pt == 12


def test_exhibit_shapes_are_not_reported_as_misaligned(decks):
    issues = run_checks(inspect_pptx(decks["drawn"]))
    flagged = {i["shape"] for i in issues if i["check"] == "misaligned"}
    assert flagged == {"Nudged note"}


def test_seeded_safe_cases_are_still_fixed(decks, tmp_path):
    for key in ("flawed1", "flawed2"):
        dst = tmp_path / f"{key}.fixed.pptx"
        applied = {(f["slide"], f["check"]) for f in fix_pptx(decks[key], dst) if f["ok"]}
        assert applied, key
    f1 = {(f["slide"], f["check"]) for f in fix_pptx(decks["flawed1"], tmp_path / "a.pptx")}
    assert {(3, "font_below_floor"), (5, "off_slide"), (6, "missing_source")} <= f1
    f2 = {(f["slide"], f["check"]) for f in fix_pptx(decks["flawed2"], tmp_path / "b.pptx")}
    assert (6, "misaligned") in f2


@pytest.mark.parametrize("deck", sorted((REPO / "examples").glob("*/deck.pptx")),
                         ids=lambda p: p.parent.name)
def test_fix_leaves_built_exhibit_shapes_alone_in_the_examples(deck, tmp_path):
    dst = tmp_path / "fixed.pptx"
    fix_pptx(deck, dst)
    assert {n for n in _changed(deck, dst) if n.startswith("sw:")} == set()


def test_fix_leaves_every_pro_showcase_shape_where_it_was(tmp_path):
    deck = _pro_showcase()
    if deck is None:
        pytest.skip("Pro showcase deck not found (set SOWHAT_DECKS_PRO)")
    dst = tmp_path / "showcase.fixed.pptx"
    fix_pptx(deck, dst)
    assert _changed(deck, dst) == set()
    issues = run_checks(inspect_pptx(deck))
    assert not [i for i in issues if i["check"] == "misaligned"]
    assert not [i for i in issues if i["check"] == "overlap" and i["severity"] == "low"
                and (i.get("shape") or "").startswith("sw:")]


@pytest.mark.parametrize("deck", sorted((REPO / "examples").glob("*/deck.pptx")),
                         ids=lambda p: p.parent.name)
def test_series_labels_are_not_reported_as_misaligned_with_the_title(deck):
    # Issue #19: a stacked bar's series names sit beside the chart, not on the
    # title's left edge.
    issues = run_checks(inspect_pptx(deck))
    assert not [i for i in issues if i["check"] == "misaligned"]
