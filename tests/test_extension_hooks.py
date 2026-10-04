"""Extension points for add-on skills: exhibit types and a template hook.

An add-on registers an exhibit type (schema branch + renderer + optional check) and the
normal validator and builder accept it. The built-in types stay as they are.
"""
import copy

import pytest
from pptx import Presentation

from deckkit import exhibits
from deckkit import spec as deckspec
from deckkit.build import build_deck
from deckkit.draw import rect
from deckkit.grid import Box

CALLOUT = {
    "type": "object",
    "required": ["type", "title", "label"],
    "additionalProperties": False,
    "properties": {
        "type": {"const": "test_callout"},
        "title": {"$ref": "#/$defs/ex_common/title"},
        "unit": {"$ref": "#/$defs/ex_common/unit"},
        "label": {"type": "string", "minLength": 1, "maxLength": 20},
        "size": {"type": "number", "minimum": 0},
    },
}


def draw_callout(ctx, slide, ex, box):
    rect(slide, Box(box.x, box.y, box.w // 2, box.h // 2), fill=ctx.accent, name="sw:test-callout")


def check_callout(ex, path, errors):
    if ex.get("size", 0) > 10:
        errors.append(f"{path}.size: callouts are at most 10")


@pytest.fixture(autouse=True, scope="module")
def callout_type():
    """Register the test type for this module only, then remove it so later tests see the stock schema."""
    exhibits.register(CALLOUT, draw_callout, check_callout)
    yield
    branches = deckspec.schema()["$defs"]["exhibit"]["oneOf"]
    branches[:] = [b for b in branches if b.get("properties", {}).get("type", {}).get("const") != "test_callout"]
    deckspec._ADDED_EXHIBITS.pop("test_callout", None)
    exhibits.RENDERERS.pop("test_callout", None)


def deck(exhibit):
    return {"spec_version": "1.0", "meta": {"title": "Hooks", "sample_data": True},
            "slides": [{"layout": "exhibit", "title": "The add-on exhibit draws where a built-in one would",
                        "exhibit": exhibit, "source": "Sample data"}]}


def names(slide):
    return [sh.name for sh in slide.shapes]


def test_registered_exhibit_validates_and_builds(tmp_path):
    prs, report = build_deck(deck({"type": "test_callout", "title": "Callout", "label": "A"}))
    prs.save(tmp_path / "deck.pptx")
    slide = Presentation(tmp_path / "deck.pptx").slides[0]
    assert "sw:test-callout" in names(slide)
    assert "sw:exhibit-title" in names(slide)  # the shared exhibit header still applies
    assert report.overflows == []


def test_registered_schema_is_enforced():
    errors = deckspec.validate(deck({"type": "test_callout", "title": "Callout", "label": "A", "colour": "red"}))
    assert any("colour: unknown field" in e for e in errors)
    errors = deckspec.validate(deck({"type": "test_callout", "title": "Callout"}))
    assert any(e.endswith(".label: required") for e in errors)


def test_registered_check_runs_after_the_schema():
    errors = deckspec.validate(deck({"type": "test_callout", "title": "Callout", "label": "A", "size": 11}))
    assert errors == ["slides[0].exhibit.size: callouts are at most 10"]


def test_unknown_types_are_still_rejected():
    errors = deckspec.validate(deck({"type": "not_registered", "title": "X"}))
    assert len(errors) == 1 and "'not_registered' is not one of" in errors[0]
    assert "test_callout" in errors[0] and "bar" in errors[0]


def test_registering_again_replaces_the_type():
    stricter = copy.deepcopy(CALLOUT)
    stricter["properties"]["label"]["maxLength"] = 2
    exhibits.register(stricter, draw_callout, check_callout)
    try:
        errors = deckspec.validate(deck({"type": "test_callout", "title": "Callout", "label": "ABC"}))
        assert errors == ["slides[0].exhibit.label: 3 characters, max 2"]
    finally:
        exhibits.register(CALLOUT, draw_callout, check_callout)
    assert deckspec.validate(deck({"type": "test_callout", "title": "Callout", "label": "ABC"})) == []


def test_built_in_types_cannot_be_replaced():
    hijack = copy.deepcopy(CALLOUT)
    hijack["properties"]["type"] = {"const": "bar"}
    with pytest.raises(ValueError, match="built-in"):
        exhibits.register(hijack, draw_callout)
    assert deckspec.validate(deck({"type": "bar", "title": "Bars", "categories": ["A", "B"], "values": [1, 2]})) == []


def test_a_schema_without_a_type_name_is_refused():
    with pytest.raises(ValueError, match="properties.type.const"):
        exhibits.register({"type": "object", "properties": {}}, draw_callout)


def test_on_template_runs_before_slides_and_can_narrow_the_frame():
    seen = []

    def narrow(tpl):
        seen.append(len(tpl.prs.slides))
        f = tpl.frame
        f.body = Box(f.body.x, f.body.y, f.body.w // 2, f.body.h)

    spec = deck({"type": "test_callout", "title": "Callout", "label": "A"})
    wide, _ = build_deck(copy.deepcopy(spec))
    narrow_prs, _ = build_deck(copy.deepcopy(spec), on_template=narrow)
    assert seen == [0]

    def callout_width(prs):
        return next(sh.width for sh in prs.slides[0].shapes if sh.name == "sw:test-callout")

    assert callout_width(narrow_prs) < callout_width(wide) * 0.6

