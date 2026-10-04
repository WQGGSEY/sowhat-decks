"""scripts/build.py as a user runs it."""
import json
import pathlib
import subprocess
import sys

from pptx import Presentation

ROOT = pathlib.Path(__file__).resolve().parents[1]
BUILD = ROOT / "skills" / "deck-build" / "scripts" / "build.py"
EXHIBITS_BUILD = ROOT / "skills" / "deck-exhibits" / "scripts" / "build.py"
FIXTURE = ROOT / "tests" / "fixtures" / "05-exhibits-bar.json"


def run(*args, cwd=None):
    return subprocess.run([sys.executable, *map(str, args)], capture_output=True, text=True, cwd=cwd, timeout=300)


def test_build_writes_deck_and_report(tmp_path):
    res = run(BUILD, FIXTURE, "-o", tmp_path, "--no-render")
    assert res.returncode == 0, res.stderr
    assert Presentation(tmp_path / "deck.pptx").slides
    report = json.loads((tmp_path / "build-report.json").read_text())
    assert report["overflow_count"] == 0 and report["checks"] == []
    assert "checks passed" in res.stdout


def test_invalid_spec_exits_2_with_the_problem(tmp_path):
    spec = json.loads(FIXTURE.read_text())
    del spec["slides"][0]["title"]
    bad = tmp_path / "bad.json"
    bad.write_text(json.dumps(spec))
    res = run(BUILD, bad, "-o", tmp_path, "--no-render")
    assert res.returncode == 2
    assert "slides[0].title: required" in res.stderr
    assert not (tmp_path / "deck.pptx").exists()


def test_strict_mode_fails_when_text_overflows(tmp_path):
    pillar = {"heading": "Heading", "body": "word " * 44, "bullets": ["word " * 20] * 4}
    spec = {"spec_version": "1.0", "meta": {"title": "Too much text"},
            "slides": [{"layout": "pillars", "title": "Four pillars with far too much text", "pillars": [pillar] * 4}]}
    path = tmp_path / "dense.json"
    path.write_text(json.dumps(spec))
    res = run(BUILD, path, "-o", tmp_path, "--no-render", "--strict")
    assert res.returncode == 3
    assert "overflow" in res.stdout
    assert (tmp_path / "deck.pptx").exists()  # the deck is still written so it can be inspected


def test_check_only_validates_without_building(tmp_path):
    res = run(BUILD, FIXTURE, "--check-only", cwd=tmp_path)
    assert res.returncode == 0 and "spec OK" in res.stdout
    assert not (tmp_path / "out").exists()


def test_export_template(tmp_path):
    res = run(BUILD, "--export-template", tmp_path / "default.pptx")
    assert res.returncode == 0
    assert len(Presentation(tmp_path / "default.pptx").slide_layouts) == 5


def test_template_flag_is_relative_to_the_working_directory(tmp_path):
    Presentation().save(tmp_path / "brand.pptx")
    res = run(BUILD, FIXTURE, "--template", "brand.pptx", "-o", "deck-out", "--no-render", cwd=tmp_path)
    assert res.returncode == 0, res.stderr
    assert Presentation(tmp_path / "deck-out" / "deck.pptx").slide_width == 9144000


def test_deck_exhibits_skill_builds_on_its_own(tmp_path):
    """npx skills add may install deck-exhibits alone: its own copy of deckkit must work."""
    res = run(EXHIBITS_BUILD, FIXTURE, "-o", tmp_path, "--no-render")
    assert res.returncode == 0, res.stderr
    assert (tmp_path / "deck.pptx").exists()
