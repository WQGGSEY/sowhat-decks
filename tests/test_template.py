"""User templates (.pptx / .potx): their layouts, placeholders, size and theme are used."""
import copy
import io
import json
import pathlib
import zipfile

import pytest
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from deckkit import checks
from deckkit.build import build_deck
from deckkit.template_map import TemplateError, open_template
from deckkit.theme import theme_fonts

ROOT = pathlib.Path(__file__).resolve().parents[1]
BOARD = json.loads((ROOT / "tests" / "fixtures" / "09-board-update-en.json").read_text(encoding="utf-8"))


@pytest.fixture()
def office_template(tmp_path):
    """python-pptx's stock 4:3 'Office Theme' template, with one sample slide left in it."""
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = "Sample slide that must not survive"
    path = tmp_path / "office.pptx"
    prs.save(path)
    return path


def as_potx(src: pathlib.Path, dst: pathlib.Path) -> pathlib.Path:
    data = io.BytesIO()
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(data, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            blob = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                blob = blob.replace(b"presentationml.presentation.main+xml", b"presentationml.template.main+xml")
            zout.writestr(item, blob)
    dst.write_bytes(data.getvalue())
    return dst


def test_deck_on_user_template_keeps_its_size_and_layouts(office_template):
    prs, report = build_deck(copy.deepcopy(BOARD), template=office_template)
    assert (prs.slide_width, prs.slide_height) == (9144000, 6858000)  # 4:3 kept
    assert len(prs.slides) == len(BOARD["slides"])  # template's sample slide dropped
    assert prs.slides[0].slide_layout.name == "Title Slide"
    assert {s.slide_layout.name for s in list(prs.slides)[1:]} == {"Title Only"}
    assert report.overflows == []


def test_deck_on_user_template_passes_the_checks(office_template, tmp_path):
    prs, _ = build_deck(copy.deepcopy(BOARD), template=office_template)
    prs.save(tmp_path / "deck.pptx")
    assert checks.inspect(Presentation(tmp_path / "deck.pptx")) == []


def test_text_is_measured_with_the_template_fonts(office_template):
    tpl = open_template(office_template)
    assert (tpl.frame.head_font, tpl.frame.body_font) == ("Calibri", "Calibri")
    assert tpl.accent == "4F81BD"


def test_potx_templates_are_accepted(office_template, tmp_path):
    potx = as_potx(office_template, tmp_path / "brand.potx")
    prs, _ = build_deck(copy.deepcopy(BOARD), template=potx)
    out = tmp_path / "deck.pptx"
    prs.save(out)
    assert len(Presentation(out).slides) == len(BOARD["slides"])


def test_layout_map_picks_named_layouts(office_template):
    spec = copy.deepcopy(BOARD)
    spec["template"] = {"layout_map": {"content": "Title and Content"}}
    prs, _ = build_deck(spec, template=office_template)
    assert prs.slides[1].slide_layout.name == "Title and Content"
    # the unused body placeholder is removed so nothing empty is left on the slide
    assert not any(ph.placeholder_format.type == PP_PLACEHOLDER.OBJECT for ph in prs.slides[1].placeholders)


def test_unknown_layout_name_lists_the_real_ones(office_template):
    with pytest.raises(TemplateError, match="Title Only"):
        open_template(office_template, layout_map={"content": "Does not exist"})


def test_spec_template_path_is_relative_to_the_spec(office_template):
    spec = copy.deepcopy(BOARD)
    spec["template"] = {"path": office_template.name}
    prs, _ = build_deck(spec, base_dir=office_template.parent)
    assert prs.slide_width == 9144000


def test_not_a_powerpoint_file_is_a_clear_error(tmp_path):
    bad = tmp_path / "notes.pptx"
    bad.write_text("not a zip")
    with pytest.raises(TemplateError, match="not a PowerPoint file"):
        open_template(bad)


def test_korean_deck_fills_an_empty_east_asian_theme_font(office_template):
    tpl = open_template(office_template, language="ko")
    assert theme_fonts(tpl.prs)["minor_ea"] == "Malgun Gothic"


def test_default_template_has_the_documented_layouts():
    tpl = open_template(None)
    assert [l.name for l in tpl.prs.slide_layouts] == ["Title Slide", "Title and Content", "Section Header",
                                                       "Title Only", "Blank"]
    assert (tpl.frame.head_font, tpl.frame.body_font) == ("Georgia", "Arial")
