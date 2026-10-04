"""The 10 committed test specs: build, re-open, and hold the quality bar.

Quality bar (c02-definition 6.6): re-opens with python-pptx, nothing off the slide,
zero estimated overflow, body >= 12 pt, source >= 10 pt, titles <= 2 lines at >= 24 pt,
exhibits are native charts or shapes (never pictures).
"""
import json
import pathlib

import pytest
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from deckkit import checks, text_fit
from deckkit.build import build_deck
from deckkit.theme import theme_fonts

ROOT = pathlib.Path(__file__).resolve().parents[1]
OUT = ROOT / "out" / "tests"
SPECS = sorted((ROOT / "tests" / "fixtures").glob("*.json"))
SMALL_ROLES = {"sw:source", "sw:page", "sw:tracker", "sw:sticker", "sw:footer"}  # 10 pt floor
CHART_TYPES = {
    ("bar", "horizontal"): XL_CHART_TYPE.BAR_CLUSTERED,
    ("bar", "vertical"): XL_CHART_TYPE.COLUMN_CLUSTERED,
    ("line", None): XL_CHART_TYPE.LINE,
    ("stacked_bar", None): XL_CHART_TYPE.COLUMN_STACKED,
    ("stacked_bar_100", None): XL_CHART_TYPE.COLUMN_STACKED_100,
}


@pytest.fixture(scope="module")
def built():
    """Build every fixture once, save it, and re-open the saved file."""
    OUT.mkdir(parents=True, exist_ok=True)
    decks = {}
    for path in SPECS:
        spec = json.loads(path.read_text(encoding="utf-8"))
        prs, report = build_deck(spec)
        target = OUT / f"{path.stem}.pptx"
        prs.save(target)
        decks[path.stem] = (spec, report, Presentation(target))
    return decks


def all_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from all_shapes(sh.shapes)
        else:
            yield sh


def run_sizes(shape):
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if r.text.strip() and r.font.size is not None:
                yield r.font.size.pt


def test_there_are_ten_committed_specs():
    assert len(SPECS) == 10


def test_specs_mark_their_numbers_as_sample_data():
    for path in SPECS:
        spec = json.loads(path.read_text(encoding="utf-8"))
        if spec.get("mode") == "ghost":
            continue
        assert spec["meta"].get("sample_data") is True, path.name


@pytest.mark.parametrize("name", [p.stem for p in SPECS])
def test_deck_reopens_with_every_slide(built, name):
    spec, _, prs = built[name]
    assert len(prs.slides) == len(spec["slides"])


@pytest.mark.parametrize("name", [p.stem for p in SPECS])
def test_no_shape_outside_the_slide(built, name):
    _, _, prs = built[name]
    w, h = prs.slide_width, prs.slide_height
    for no, slide in enumerate(prs.slides, 1):
        for sh in all_shapes(slide.shapes):
            if sh.left is None:
                continue
            assert sh.left >= 0 and sh.top >= 0, f"slide {no} {sh.name}"
            assert sh.left + sh.width <= w and sh.top + sh.height <= h, f"slide {no} {sh.name}"


@pytest.mark.parametrize("name", [p.stem for p in SPECS])
def test_zero_estimated_overflow(built, name):
    _, report, prs = built[name]
    assert report.overflows == []
    assert [i for i in checks.inspect(prs) if i.kind == "overflow"] == []


@pytest.mark.parametrize("name", [p.stem for p in SPECS])
def test_body_at_least_12pt_and_source_at_least_10pt(built, name):
    _, _, prs = built[name]
    for no, slide in enumerate(prs.slides, 1):
        for sh in all_shapes(slide.shapes):
            if sh.has_text_frame:
                floor = 10 if sh.name in SMALL_ROLES else 12
                for size in run_sizes(sh):
                    assert size >= floor, f"slide {no} {sh.name} {size}pt"
            if sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        for size in run_sizes(cell):
                            assert size >= 12, f"slide {no} table {size}pt"
            if sh.has_chart:
                for el in sh.chart._chartSpace.iter(qn("a:defRPr"), qn("a:rPr")):
                    if el.get("sz"):
                        assert int(el.get("sz")) >= 1200, f"slide {no} chart text {el.get('sz')}"


@pytest.mark.parametrize("name", [p.stem for p in SPECS])
def test_titles_fit_two_lines_at_24pt_or_more(built, name):
    spec, _, prs = built[name]
    for no, (slide, s) in enumerate(zip(prs.slides, spec["slides"]), 1):
        title = slide.shapes.title
        if s["layout"] in ("cover", "section"):
            continue
        sizes = list(run_sizes(title))
        assert sizes and min(sizes) >= 24, f"slide {no}"
        width = (title.width - title.text_frame.margin_left - title.text_frame.margin_right) / 12700
        lines = text_fit.wrap(title.text_frame.text, "Georgia", min(sizes), width)
        assert len(lines) <= 2, f"slide {no}: {lines}"


@pytest.mark.parametrize("name", [p.stem for p in SPECS])
def test_exhibits_are_native_objects(built, name):
    spec, _, prs = built[name]
    for no, (slide, s) in enumerate(zip(prs.slides, spec["slides"]), 1):
        shapes = list(all_shapes(slide.shapes))
        assert not [sh for sh in shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE], f"slide {no} has a picture"
        ex = s.get("exhibit")
        if not ex or spec.get("mode") == "ghost" or s["layout"] == "ghost":
            continue
        kind = ex["type"]
        if kind in ("bar", "line", "stacked_bar", "stacked_bar_100"):
            orient = ex.get("orientation", "horizontal") if kind == "bar" else None
            charts = [sh.chart for sh in shapes if sh.has_chart]
            assert [c.chart_type for c in charts] == [CHART_TYPES[(kind, orient)]], f"slide {no}"
            assert not charts[0].has_legend, f"slide {no}: use direct labels, not a legend"
        elif kind == "highlight_table":
            assert any(sh.has_table for sh in shapes), f"slide {no}"
        elif kind == "matrix_2x2":
            assert sum(sh.name == "sw:quadrant" for sh in shapes) == 4, f"slide {no}"
        elif kind == "process":
            steps = [sh for sh in slide.shapes if sh.name == "sw:step"]
            assert len(steps) == len(ex["steps"]), f"slide {no}"


def test_highlighted_bar_uses_the_accent_and_others_are_grey(built):
    _, _, prs = built["05-exhibits-bar"]
    chart = next(sh.chart for sh in prs.slides[0].shapes if sh.has_chart)
    series = chart.plots[0].series[0]
    assert str(series.format.fill.fore_color.rgb) == "B7BEC7"
    # sorted descending, so the highlighted Product 4 is the first bar
    assert str(series.points[0].format.fill.fore_color.rgb) == "1F5AA6"


def test_line_chart_end_labels_name_the_series(built):
    _, _, prs = built["06-exhibits-line"]
    chart = next(sh.chart for sh in prs.slides[1].shapes if sh.has_chart)
    labels = [s._element.find(qn("c:dLbls")) for s in chart.plots[0].series]
    texts = ["".join(t.text for t in d.iter(qn("a:t"))) for d in labels]
    assert texts == ["Channel A 24", "Channel B 23", "Channel C 34", "Channel D 19"]


def test_sample_data_is_marked_on_the_slide(built):
    _, _, prs = built["01-en-all-layouts"]
    sources = [sh.text_frame.text for sh in prs.slides[3].shapes if sh.name == "sw:source"]
    assert sources and "Sample data" in sources[0]


def test_draft_sticker_and_page_numbers(built):
    _, _, prs = built["01-en-all-layouts"]
    for no, slide in enumerate(prs.slides, 1):
        names = [sh.name for sh in slide.shapes]
        assert "sw:sticker" in names
        assert ("sw:page" in names) == (no != 1)  # no page number on the cover


def test_ghost_mode_shows_only_titles_and_placeholder_boxes(built):
    spec, _, prs = built["04-ghost-en"]
    for slide, s in zip(prs.slides, spec["slides"]):
        names = {sh.name for sh in slide.shapes}
        assert "sw:sticker" in names
        if s["layout"] in ("cover", "section"):
            continue
        assert not any(sh.has_chart for sh in slide.shapes)
        assert "sw:ghost-box" in names


@pytest.mark.parametrize("name,font,tag", [("02-ko-all-layouts", "Malgun Gothic", "ko-KR"),
                                           ("03-ja-smoke", "Yu Gothic", "ja-JP")])
def test_east_asian_decks_set_theme_font_and_language(built, name, font, tag):
    _, _, prs = built[name]
    assert theme_fonts(prs)["minor_ea"] == font
    assert theme_fonts(prs)["major_ea"] == font
    title_run = prs.slides[1].shapes.title.text_frame.paragraphs[0].runs[0]
    assert title_run._r.get_or_add_rPr().get("lang") == tag


def test_speaker_notes_survive_the_round_trip(built):
    _, _, prs = built["01-en-all-layouts"]
    assert prs.slides[0].notes_slide.notes_text_frame.text.startswith("Open with the answer")


@pytest.mark.parametrize("name", [p.stem for p in SPECS])
def test_drawn_shapes_carry_no_theme_style(built, name):
    """A theme style reference brings the theme's shadow along (it did on rules and 2x2 arrows)."""
    _, _, prs = built[name]
    for no, slide in enumerate(prs.slides, 1):
        for el in slide.shapes._spTree.iter(qn("p:sp"), qn("p:cxnSp")):
            if el.find(".//" + qn("p:ph")) is None:
                assert el.find(qn("p:style")) is None, f"slide {no}"
