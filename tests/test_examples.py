"""The published examples stay what README.md says they are.

Every example folder has its files, a valid spec, a deck that matches the spec,
native charts with source lines, and no high or medium deck-review findings.
The draft deck in example 03 keeps the problems it was built to show.
"""
import importlib.util
import pathlib
import subprocess
import sys

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = pathlib.Path(__file__).resolve().parents[1]
REVIEW_SCRIPTS = ROOT / "skills" / "deck-review" / "scripts"
if str(REVIEW_SCRIPTS) not in sys.path:
    sys.path.insert(0, str(REVIEW_SCRIPTS))

from deckkit import spec as deckspec  # noqa: E402  (tests/conftest.py puts deck-build's scripts on the path)
from deckreview import checks as review_checks  # noqa: E402
from deckreview.inspect import inspect_pptx  # noqa: E402

EXAMPLES = ["01-investor-update", "02-market-entry", "03-review-before-after"]
FILES = ["README.md", "brief.md", "storyline.md", "deck.json", "deck.pptx", "deck.pdf",
         "review.md", "SOURCES.md", "preview/slide-01.png"]
STORYLINE_TOOL = ROOT / "skills" / "deck-storyline" / "scripts" / "storyline_tool.py"


def example(name):
    return ROOT / "examples" / name


@pytest.mark.parametrize("name", EXAMPLES)
def test_example_has_every_file(name):
    missing = [f for f in FILES if not (example(name) / f).is_file()]
    assert missing == []


@pytest.mark.parametrize("name", EXAMPLES)
def test_storyline_passes_its_check(name):
    proc = subprocess.run([sys.executable, str(STORYLINE_TOOL), "check", str(example(name) / "storyline.md")],
                          capture_output=True, text=True, check=False)
    assert proc.returncode == 0, proc.stdout + proc.stderr


@pytest.mark.parametrize("name", EXAMPLES)
def test_deck_matches_its_spec_and_previews(name):
    spec = deckspec.load(example(name) / "deck.json")
    prs = Presentation(str(example(name) / "deck.pptx"))
    assert len(prs.slides) == len(spec["slides"])
    assert len(list((example(name) / "preview").glob("slide-*.png"))) == len(spec["slides"])


@pytest.mark.parametrize("name", EXAMPLES)
def test_charts_are_native_and_every_exhibit_has_a_source(name):
    prs = Presentation(str(example(name) / "deck.pptx"))
    charts = 0
    for number, slide in enumerate(prs.slides, start=1):
        shapes = list(slide.shapes)
        assert not [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE], f"picture on slide {number}"
        text = " ".join(s.text_frame.text for s in shapes if s.has_text_frame)
        if any(s.has_chart or s.has_table for s in shapes):
            assert "Source:" in text, f"slide {number} has an exhibit without a source line"
        charts += sum(1 for s in shapes if s.has_chart)
    assert charts >= 4


@pytest.mark.parametrize("name", EXAMPLES)
def test_deck_review_finds_nothing_high_or_medium(name):
    issues = review_checks.run_checks(inspect_pptx(example(name) / "deck.pptx"))
    serious = [(i["slide"], i["check"], i["severity"]) for i in issues if i["severity"] in ("high", "medium")]
    assert serious == []


def test_sample_data_example_says_so_on_every_slide():
    prs = Presentation(str(example("03-review-before-after") / "deck.pptx"))
    for number, slide in enumerate(prs.slides, start=1):
        text = " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame).lower()
        assert "sample data" in text, f"slide {number}"


def test_before_deck_keeps_its_planted_problems():
    issues = review_checks.run_checks(inspect_pptx(example("03-review-before-after") / "before" / "before.pptx"))
    found = {i["check"] for i in issues}
    expected = {"title_not_claim", "font_below_floor", "missing_source", "data_needed"}
    if importlib.util.find_spec("PIL") is not None:  # the picture-of-a-chart check reads the image with Pillow
        expected.add("chart_as_image")
    assert expected <= found
