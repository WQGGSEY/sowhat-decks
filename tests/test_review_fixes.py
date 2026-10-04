"""Regression tests for the PR #1 review (B1-B5)."""
import ast
import copy
import json
import math
import pathlib
import shutil
import subprocess
import sys

import pytest
from pptx import Presentation

from deckkit import checks, render, spec as deckspec
from deckkit.build import build_deck
from deckkit.draw import source_text

ROOT = pathlib.Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "skills" / "deck-build" / "scripts"


def deck(*slides, **meta):
    return {"spec_version": "1.0", "meta": {"title": "Review fixes", **meta}, "slides": list(slides)}


def bar_slide(**exhibit):
    ex = {"type": "bar", "title": "Growth by region", "categories": ["A", "B"], "values": [1, 2]}
    ex.update(exhibit)
    return {"layout": "exhibit", "title": "Region B grew fastest", "exhibit": ex, "source": "Finance"}


def texts(slide, name):
    return [sh.text_frame.text for sh in slide.shapes if sh.name == name]


# ---------------------------------------------------------------- B1: Python 3.9

def _python39():
    for cand in ("python3.9", "/usr/bin/python3"):
        exe = shutil.which(cand) or (cand if pathlib.Path(cand).is_file() else None)
        if exe:
            out = subprocess.run([exe, "-c", "import sys; print(sys.version_info[:2] == (3, 9))"],
                                 capture_output=True, text=True)
            if out.stdout.strip() == "True":
                return exe
    return None


def test_library_parses_as_python_39():
    for path in list((SCRIPTS / "deckkit").glob("*.py")) + list(SCRIPTS.glob("*.py")) + list((ROOT / "tools").glob("*.py")):
        ast.parse(path.read_text(encoding="utf-8"), filename=str(path), feature_version=(3, 9))


def test_no_python_310_only_keyword_arguments():
    # stdlib keywords added in 3.10+: TemporaryDirectory(ignore_cleanup_errors=), zip(strict=), dataclass(kw_only=, slots=)
    banned = {"ignore_cleanup_errors", "strict", "kw_only", "slots"}
    for path in (SCRIPTS / "deckkit").glob("*.py"):
        tree = ast.parse(path.read_text(encoding="utf-8"))
        used = {kw.arg for node in ast.walk(tree) if isinstance(node, ast.Call) for kw in node.keywords}
        assert not used & banned, f"{path.name}: {used & banned}"


@pytest.mark.skipif(_python39() is None or render.find_soffice() is None, reason="needs python3.9 and LibreOffice")
def test_render_works_on_python_39(tmp_path):
    prs, _ = build_deck(deck(bar_slide()))
    prs.save(tmp_path / "deck.pptx")
    code = (f"import sys; sys.path.insert(0, {str(SCRIPTS)!r}); from deckkit import render; "
            f"r = render.render({str(tmp_path / 'deck.pptx')!r}, {str(tmp_path)!r}); print(r.pdf)")
    res = subprocess.run([_python39(), "-c", code], capture_output=True, text=True, timeout=300)
    assert res.returncode == 0, res.stderr
    assert (tmp_path / "deck.pdf").stat().st_size > 1000


# ---------------------------------------------------------------- B2: sample-data label

class _Ctx:
    def __init__(self, lang="en", sample=False):
        from deckkit.theme import LABELS
        self.labels = LABELS[lang]
        self.meta = {"language": lang, "sample_data": sample}


@pytest.mark.parametrize("source,sample,lang,expected", [
    ("Customer survey, sample of 500 CFOs", True, "en", "Source: Sample data; Customer survey, sample of 500 CFOs"),
    ("Source: Finance export", True, "en", "Source: Sample data; Finance export"),
    ("Source: Finance export", False, "en", "Source: Finance export"),
    ("source: finance export", False, "en", "source: finance export"),
    ("Sources: A; B", True, "en", "Sources: Sample data; A; B"),
    ("Sample data", True, "en", "Source: Sample data"),
    ("Sample data", False, "en", "Source: Sample data"),
    ("", True, "en", "Source: Sample data"),
    ("", False, "en", ""),
    ("Sample data; finance export", True, "en", "Source: Sample data; finance export"),
    ("Survey of sample stores", False, "en", "Source: Survey of sample stores"),
    ("출처: 고객 설문", True, "ko", "출처: 샘플 데이터; 고객 설문"),
    ("샘플 데이터", True, "ko", "출처: 샘플 데이터"),
    ("出所：社内データ", False, "ja", "出所：社内データ"),
])
def test_source_line_labels_are_explicit(source, sample, lang, expected):
    assert source_text(_Ctx(lang, sample), {"source": source}) == expected


# ---------------------------------------------------------------- B3: ghost evidence

@pytest.mark.parametrize("item", ["[DATA NEEDED: partner terms]", "[DATA NEEDED] partner terms",
                                  "partner terms [data needed]", "[ DATA NEEDED : partner terms ]"])
def test_ghost_puts_every_data_needed_form_under_still_needed(item):
    prs, _ = build_deck(deck({"layout": "ghost", "title": "Enter through a partner",
                              "ghost": {"evidence": ["E1 Market size [SRC 1]", item]}}))
    box = next(t for t in texts(prs.slides[0], "sw:body") if "partner terms" in t.lower())
    lines = box.split("\n")
    assert lines.index("Still needed") < lines.index(item)
    assert lines.index("Evidence") < lines.index("E1 Market size [SRC 1]") < lines.index("Still needed")


# ---------------------------------------------------------------- B4: validation never crashes

WRONG = [None, [], {}, 7, "x", True, ["a", 1], {"a": 1}, float("nan")]
RICH = deck(
    {"layout": "cover"},
    bar_slide(highlight=["B"], number_format="0.0"),
    {"layout": "exhibit_takeaways", "title": "T", "takeaways": ["x"], "source": "S",
     "exhibit": {"type": "line", "title": "L", "categories": ["1", "2"], "series": [{"name": "s", "values": [1, 2]}], "highlight": "s"}},
    {"layout": "exhibit", "title": "T", "source": "S",
     "exhibit": {"type": "stacked_bar", "title": "S", "categories": ["1"], "series": [{"name": "a", "values": [1]}, {"name": "b", "values": [2]}]}},
    {"layout": "table", "title": "T", "table": {"columns": ["a", "b"], "rows": [["x", 1]], "highlight": {"rows": [0]}, "number_format": ",.1f"}},
    {"layout": "executive_summary", "title": "T", "points": ["a", {"headline": "h", "detail": "d"}]},
    {"layout": "appendix", "title": "T", "bullets": ["a"]},
    {"layout": "ghost", "title": "T", "ghost": {"evidence": ["e"]}},
)


def _paths(node, path=()):
    yield path
    if isinstance(node, dict):
        for k, v in node.items():
            yield from _paths(v, path + (k,))
    elif isinstance(node, list):
        for i, v in enumerate(node):
            yield from _paths(v, path + (i,))


def _set(node, path, value):
    for key in path[:-1]:
        node = node[key]
    node[path[-1]] = value


def test_validator_never_raises_on_wrongly_typed_fields():
    assert deckspec.validate(RICH) == []
    for path in list(_paths(RICH))[1:]:
        for value in WRONG:
            bad = copy.deepcopy(RICH)
            _set(bad, path, value)
            errors = deckspec.validate(bad)  # must not raise
            assert all(isinstance(e, str) for e in errors)


@pytest.mark.parametrize("value", [None, 3, [], "x"])
def test_validator_handles_a_non_object_spec(value):
    assert deckspec.validate(value)


def test_unhashable_discriminator_is_an_error_not_a_crash():
    errors = deckspec.validate(deck({"layout": ["cover"]}))
    assert any("layout" in e for e in errors)


@pytest.mark.parametrize("bad", [float("nan"), float("inf"), -float("inf")])
def test_non_finite_numbers_are_rejected(bad):
    errors = deckspec.validate(deck(bar_slide(values=[1, bad])))
    assert any("finite" in e for e in errors)


def test_nan_in_a_spec_file_is_a_spec_error(tmp_path):
    path = tmp_path / "deck.json"
    path.write_text(json.dumps(deck(bar_slide())).replace("[1, 2]", "[1, NaN]"))
    with pytest.raises(deckspec.SpecError, match="finite"):
        deckspec.load(path)


def test_stacked_bars_reject_negative_values():
    slide = {"layout": "exhibit", "title": "T", "source": "S",
             "exhibit": {"type": "stacked_bar", "title": "S", "categories": ["1", "2"],
                         "series": [{"name": "a", "values": [-1, -2]}, {"name": "b", "values": [-3, -4]}]}}
    errors = deckspec.validate(deck(slide))
    assert any("0 or more" in e for e in errors)


@pytest.mark.parametrize("fmt", ["0&0", '0"x"', "<0>", "abc"])
def test_unsafe_chart_number_formats_are_rejected(fmt):
    errors = deckspec.validate(deck(bar_slide(number_format=fmt)))
    assert any("number_format" in e for e in errors)


@pytest.mark.parametrize("fmt,label", [("$#,##0", "$1,235"), ("#,##0.0", "1,234.6"), ("0%", "123456%"), ("0.0", "1234.6")])
def test_chart_number_formats_build_and_label_alike(fmt, label):
    prs, _ = build_deck(deck(bar_slide(values=[1, 1234.56], number_format=fmt)))
    from deckkit.exhibits import fmt_number
    assert fmt_number(1234.56, fmt) == label


@pytest.mark.parametrize("fmt,shown", [("#,##0", "1,235"), ("#,##0.0", "1,234.6"), (",.2f", "1,234.56"), ("0.0", "1234.6")])
def test_table_number_format_accepts_excel_and_python_styles(fmt, shown):
    slide = {"layout": "table", "title": "T", "table": {"columns": ["a", "b"], "rows": [["x", 1234.56]], "number_format": fmt}}
    assert deckspec.validate(deck(slide)) == []
    prs, _ = build_deck(deck(slide))
    table = next(sh.table for sh in prs.slides[0].shapes if sh.has_table)
    assert table.cell(1, 1).text_frame.text == shown


def test_schema_and_library_share_the_chart_format_rule():
    from deckkit.numbers import EXCEL_FORMAT
    assert deckspec.schema()["$defs"]["ex_common"]["number_format"]["pattern"] == EXCEL_FORMAT


def test_bad_table_number_format_is_rejected():
    slide = {"layout": "table", "title": "T", "table": {"columns": ["a", "b"], "rows": [["x", 1]], "number_format": "%Q"}}
    assert any("number_format" in e for e in deckspec.validate(deck(slide)))


TRICKY = [
    bar_slide(values=[None, None]),
    bar_slide(values=[-5, -2]),
    bar_slide(values=[0, 0]),
    bar_slide(values=[1e12, 3e12]),
    bar_slide(values=[0.0001, 0.0002]),
    bar_slide(categories=["R&D <core>", "\"Ops\""], values=[1, 2], title="A & B <x>"),
    {"layout": "exhibit", "title": "T", "source": "S", "exhibit": {"type": "line", "title": "L", "categories": ["1", "2"],
                                                                     "series": [{"name": "a&b", "values": [None, None]}, {"name": "c", "values": [3, 3]}]}},
    {"layout": "exhibit", "title": "T", "source": "S", "exhibit": {"type": "line", "title": "L", "categories": ["1", "2"],
                                                                     "series": [{"name": "neg", "values": [-3, -7]}]}},
    {"layout": "exhibit", "title": "T", "source": "S", "exhibit": {"type": "stacked_bar", "title": "S", "categories": ["1", "2"],
                                                                     "series": [{"name": "a", "values": [0, 0]}, {"name": "b", "values": [None, 0]}]}},
    {"layout": "exhibit", "title": "T", "source": "S", "exhibit": {"type": "stacked_bar_100", "title": "S", "categories": ["1"],
                                                                     "series": [{"name": "a", "values": [0]}, {"name": "b", "values": [0]}]}},
]


@pytest.mark.parametrize("slide", TRICKY, ids=range(len(TRICKY)))
def test_valid_specs_always_build(slide):
    spec = deck(slide)
    assert deckspec.validate(spec) == []
    build_deck(spec)


def test_cli_reports_wrong_types_without_a_traceback(tmp_path):
    bad = deck(bar_slide())
    bad["slides"][0]["exhibit"]["categories"] = "A,B"
    path = tmp_path / "deck.json"
    path.write_text(json.dumps(bad))
    res = subprocess.run([sys.executable, str(SCRIPTS / "build.py"), str(path), "--no-render", "-o", str(tmp_path)],
                         capture_output=True, text=True)
    assert res.returncode == 2
    assert "Traceback" not in res.stderr and "categories" in res.stderr


# ---------------------------------------------------------------- B5: rows that cannot fit

def _ko(n):
    return ("매출 성장은 " * 60)[:n].strip()


@pytest.mark.parametrize("n,head,detail", [(4, 90, 180), (4, 120, 240), (5, 120, 240), (5, 90, 180)])
def test_summary_rows_stay_above_the_source_line(n, head, detail, tmp_path):
    slide = {"layout": "executive_summary", "title": "요약", "source": "샘플 데이터",
             "points": [{"headline": _ko(head), "detail": _ko(detail)} for _ in range(n)]}
    prs, report = build_deck(deck(slide, language="ko"))
    prs.save(tmp_path / "d.pptx")
    issues = checks.inspect(Presentation(tmp_path / "d.pptx"))
    body = [sh for sh in prs.slides[0].shapes if sh.name in ("sw:body", "sw:number")]
    source = next(sh for sh in prs.slides[0].shapes if sh.name == "sw:source")
    assert all(sh.top + sh.height <= source.top for sh in body)  # rows never run into the source line
    assert not any(i.kind in ("overlap", "out-of-bounds") for i in issues)
    if n == 5 and head == 120:  # five long Korean points cannot fit at 12 pt: that must be reported
        assert report.overflows and any(i.kind == "overflow" for i in issues)


def test_recommendations_that_cannot_fit_are_reported():
    items = [{"action": _ko(100), "detail": _ko(160), "owner": "담당자", "due": "2027년 1분기"}] * 6
    prs, report = build_deck(deck({"layout": "recommendations", "title": "권고", "items": items}, language="ko"))
    assert report.overflows


def test_checker_flags_text_boxes_that_overlap(tmp_path):
    prs, _ = build_deck(deck(bar_slide()))
    slide = prs.slides[0]
    source = next(sh for sh in slide.shapes if sh.name == "sw:source")
    title = slide.shapes.title
    title.top = source.top - 10000  # push the title onto the source line
    prs.save(tmp_path / "d.pptx")
    assert any(i.kind == "overlap" for i in checks.inspect(Presentation(tmp_path / "d.pptx")))
